#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera xlsx e MD das Restrições com coluna DISCIPLINA.
Arquivo: 2026-05-21_CONTROLE_CRITICO_PLA_Restricoes-Consolidadas-25098-25103_Abelv
"""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
from collections import Counter

def pf(h): return PatternFill(start_color=h, end_color=h, fill_type="solid")
def ft(bold=False, color="000000", size=10): return Font(bold=bold, color=color, size=size, name="Calibri")
def al(h="center", v="center", wrap=False): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def bt():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

C_HDR_BG="1F4E79"; C_HDR_FG="FFFFFF"
C_NP_BG="C6EFCE";  C_NP_FG="375623"
C_AL_BG="FFEB9C";  C_AL_FG="7F6000"
C_AT_BG="FFC7CE";  C_AT_FG="9C0006"
C_CO_BG="DAEEF3";  C_CO_FG="17375E"
C_CA_BG="FFF2CC";  C_CA_FG="7F4F00"
C_CI_BG="EBF3FB";  C_FA_BG="FFF2CC"

today = date(2026, 5, 21)

def status(nec, conc):
    if nec is None: return "—"
    if conc is not None: return "CONCLUÍDO" if conc <= nec else "CONCLUÍDO COM ATRASO"
    if nec < today: return "ATRASADO"
    if nec == today: return "ALERTA"
    return "NO PRAZO"

def dias(nec, conc):
    if nec is None: return "—"
    if conc is not None: return max(0, (conc - nec).days)
    if nec < today: return (today - nec).days
    return 0

# Tupla: (disc, elab, proj, emp, resp, desc, imp, id_imp, nec, obs, conc)
dados = [
    ("CIV", date(2026,5,11),"CISTO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Mobilização da Krrom para realização da Tarefa do Dique no Piso Técnico",
     None,None,date(2026,5,20),None,date(2026,5,11)),

    ("CIV", date(2026,5,8),"CISTO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Recebimento de Materiais do Dique no Piso Técnico",
     None,None,date(2026,5,11),None,date(2026,5,11)),

    ("CIV", date(2026,5,8),"CISTO","CRISTÁLIA","Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Retirada de 2 Peças de Vidro da Sala Limpa, pois está impedindo a continuidade da Marcação das Bases que serão construídas.",
     "Impede a continuação da Locação das Bases que serão construídas na Sala Limpa.",
     None,date(2026,5,19),None,None),

    ("TUB", date(2026,5,5),"CISTO","ABELV","Leobino (Eng. Civil Abelv)",
     "Chegada de Materiais para Fabricação da Tubulação do Trap",
     None,None,date(2026,5,7),"19/05: Pendente a entrega dos materiais",None),

    ("CIV", date(2026,5,14),"CISTO","CRISTÁLIA","Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Cobrar Definição da Cristália a respeito da Locação da Plataforma dos Reatores em Função da Instalação da Nova Linha Hidrossanitária.",
     "Impacta o início do Corte do Piso para confecção das Bases da Plataforma dos Reatores "
     "(Início Previsto: 21/05 - Atividade do Cronograma: ID 22). E Demolição do Piso para "
     "Instalação da Rede Hidrossanitária (Início Previsto: 18/05 - Atividade do Cronograma: ID 58).",
     None,date(2026,5,18),None,None),

    ("TUB", date(2026,5,12),"CISTO","ABELV","Tiago (Gestor de Projetos Abelv)",
     "Acompanhamento / Recebimento das Conexões de Aço Carbono",
     "Impacta na Fabricação da Tubulação de Aço Carbono "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 54 e ID 317).",
     None,date(2026,5,19),None,None),

    ("TUB", date(2026,5,12),"CISTO","ABELV","Tiago (Gestor de Projetos Abelv)",
     "Acompanhamento / Recebimento da Tubulação de Aço Carbono",
     "Impacta na Fabricação da Tubulação de Aço Carbono "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 54 e ID 317).",
     None,date(2026,5,15),None,date(2026,5,15)),

    ("HVA", date(2026,5,15),"CISTO","ABELV","Uehara (Engenheiro Mecânico do HVAC Abelv)",
     "Validar a Mudança Dimensional da Base do Lavador de Gases",
     "Impacta no Início da Demolição do Piso do Radier do Lavador de Gases a Título de Área "
     "caso aconteça alguma mudança significativa. "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 466).",
     None,date(2026,5,20),None,None),

    ("HVA", date(2026,5,12),"CISTO","ABELV","Uehara (Engenheiro Mecânico do HVAC Abelv)",
     "Contratação Lavador de Gases",
     None,None,date(2026,5,15),
     "19/05: Aguardando o Setor de Suprimentos da Abelv Fechar o Pedido.",None),

    ("SUP", date(2026,5,19),"CISTO","ABELV","Ana Assis (Gerente do Suprimentos Abelv)",
     "Fechamento do Pedido de Compra do Lavador de Gases",
     "Impacta na identificação real do Dimensional da Base do Lavador de Gases, "
     "consequentemente também impactando o Início da Demolição do Piso do Radier do Lavador de Gases "
     "a Título de Área caso aconteça alguma mudança significativa. "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 466).",
     None,date(2026,5,19),None,date(2026,5,20)),

    ("CIV", date(2026,5,12),"FARMO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Definição do Tipo de Fundação do Prédio",
     "Impacta no início da Fundação do Prédio "
     "(Início Previsto: 20/08/26 - Atividade do Cronograma: ID 431).",
     431,date(2026,5,13),None,None),

    ("CIV", date(2026,5,12),"FARMO","CRISTÁLIA","Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Recebimento do Projeto da Base do Lavador de Gases Existente.",
     "Impacta na Formalização da Nova Locação que foi combinada IN LOCO para execução da base do "
     "Lavador existente. Início da Confecção da Nova Base (Radier) do Lavador de Gases Existente "
     "(Início Previsto: 05/06/26 - Atividade do Cronograma: ID 251).",
     None,date(2026,5,15),None,None),

    ("MEC", date(2026,5,7),"FARMO","MGC","Yohanna (Engenheira Mecânica Abelv)",
     "Solicitação e Recebimento dos Projetos do Farmoquímico da Empresa MGC, e envio para Cristália para aprovação",
     "Impacta na Compra dos Materiais e Subcontratações para realização dos materiais.",
     None,date(2026,5,19),None,date(2026,5,20)),

    ("CIV", date(2026,5,19),"FARMO","ABELV","Leobino (Eng. Civil Abelv)",
     "Enviar para a MGC as cotas de elevação do Lavador de Gases",
     "Impacta na confecção do Projeto do Radier do Lavador de Gases",
     None,date(2026,5,19),None,None),

    ("ELE", date(2026,5,19),"FARMO","ABELV","Leobino (Eng. Civil Abelv)",
     "Enviar para a MGC as cotas de elevação da Sala Elétrica",
     "Impacta na confecção do Projeto da Sala Elétrica",
     None,date(2026,5,20),None,None),

    ("CIV", date(2026,5,19),"FARMO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "A Krrom precisa informar a Data em que irão apresentar o Tipo de Fundação que será aplicada no Farmoquímico.",
     "Impacta na Contratação da Empresa para realização da Estaca, por ventura impactando na compra "
     "de materiais da estaca (Início Previsto: 08/06 - Atividade do Cronograma: ID 118).",
     None,date(2026,5,19),None,None),

    ("CIV", date(2026,5,19),"CISTO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "A Krrom precisa apresentar a Metodologia de Corte que será aplicado para avaliar a resistência "
     "do piso onde poderá ser construída a base da Plataforma dos Reatores, em função da Interferência "
     "da Instalação da Nova Rede Hidrossanitária.",
     "Impacta o início do Corte do Piso para confecção das Bases da Plataforma dos Reatores "
     "(Início Previsto: 21/05 - Atividade do Cronograma: ID 22). E Demolição do Piso para "
     "Instalação da Rede Hidrossanitária (Início Previsto: 18/05 - Atividade do Cronograma: ID 58).",
     None,date(2026,5,20),
     "20/05: A Krrom apresentou a alternativa de realização de um Ensaio Não Destrutivo de Esclerometria, "
     "pois seria mais rápido o método de avaliação da Resistência Mecânica.",
     date(2026,5,20)),

    ("CIV", date(2026,5,19),"FARMO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Elaboração do Pedido da Tubulação de Água Pluvial",
     "Impacta a Compra da Tubulação de Água Pluvial",
     None,date(2026,5,27),None,None),

    ("MED", date(2026,5,19),"FARMO","KROOM","Rodrigo (Téc. Planej. da Kroom) / Leobino (Eng. Civil Abelv)",
     "A Kroom precisa enviar os Pedidos de Faturamento Direto",
     "Impacto na Medição",
     None,date(2026,5,19),None,date(2026,5,19)),

    ("SUP", date(2026,5,19),"CISTO","KROOM","Tiago (Gestor de Projetos Abelv) / Celso (Gestor de Projetos Abelv)",
     "Verificação do Efetivo de Andaime para atendimento das Obras (Farmoquímico e Citostático).",
     "Impacto de acesso para realização das Frentes de Montagem",
     None,date(2026,5,20),None,date(2026,5,20)),

    ("TUB", date(2026,5,19),"CISTO","ABELV","Leobino (Eng. Civil Abelv)",
     "Montagem do Pipe-Shop de Aço Carbono no Canteiro de Obra Abelv.",
     "Impacta na Fabricação da Tubulação de Aço Carbono "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 54 e ID 317).",
     None,date(2026,5,20),None,date(2026,5,20)),

    ("GER", date(2026,5,19),"CISTO","ABELV","Erik Massola (Comprador Abelv) / Tiago Tiago (Gestor de Projetos Abelv)",
     "O Setor de Compras precisa enviar o Efetivo (Nomes e Cpfs dos colaboradores) da empresa "
     "contratada Açoplast (Estrutura Metálica) para Tiago Matos, para posterior envio para Cristália "
     "realizar a avaliação social dos mesmos.",
     "Impacto: Caso os mesmos não forem aprovados, eles não poderão adentrar ao site, podendo "
     "impactar o prazo da obra.",
     None,date(2026,5,20),None,None),

    ("SUP", date(2026,5,19),"CISTO","ABELV","Edimar Cunha (Supervisor de Suprimentos Abelv) / Tiago Tiago (Gestor de Projetos Abelv)",
     "Contratação do Fornecedor de Andaime.",
     "Impacto de acesso para realização das Frentes de Montagem",
     None,date(2026,5,22),None,None),

    ("JUR", date(2026,5,19),"FARMO","ABELV","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "A empresa Krrom precisa Regularizar a Documentação da Empresa Recicla (Bota Fora).",
     "Impacto: Caso aconteça alguma auditoria ou acidente, estamos descobertos por documentação "
     "legal de subcontratação.",
     None,date(2026,5,21),None,None),

    ("MET", date(2026,5,19),"CISTO","ABELV","Yohanna (Engenheira Mecânica Abelv)",
     "Verificação da readequação do Projeto da Plataforma dos Reatores por base do nosso fornecedor "
     "que vai fabricar, pois os reatores adquiridos pela Cristália não cabem no mesmo.",
     "Impacto: Atraso na Fabricação do Projeto da Plataforma dos Reatores, por necessidade de "
     "elaboração do Novo Projeto para readequação do novo lay-out "
     "(Início e Término Previsto: 14/05/26 e 13/07/26 - Atividade do Cronograma: ID 650).",
     650,date(2026,5,20),None,None),

    ("MEC", date(2026,5,19),"CISTO","CRISTÁLIA","Yohanna (Engenheira Mecânica Abelv)",
     "O Filtro Secador adquirido pela Cristália não cabe na base do Projeto Fornecido, logo a Cristália "
     "ficou de elaborar um Novo Projeto da Base do Filtro Secador para essa nova realidade.",
     "Impacto: Atraso na Construção da Nova Base do Filtro Secador "
     "(Início Previsto: 20/05/26 - Atividade do Cronograma: ID 9).",
     9,date(2026,5,20),None,None),

    ("CIV", date(2026,5,19),"CISTO","CRISTÁLIA","Yohanna (Engenheira Mecânica Abelv)",
     "A Cristália tem que apresentar o Projeto da Nova Locação do Ralo do Sistema Hidrossanitário, "
     "pois o Projeto Atual da Base da Plataforma dos Reatores está na mesma coordenada dos Ralos.",
     "Impacto: Atraso na Conclusão do Corte do Piso para Novo Sistema Hidrossanitário. "
     "(Término Previsto: 18/05/26 - Atividade do Cronograma: ID 59).",
     59,date(2026,5,20),None,None),

    ("CIV", date(2026,5,20),"CISTO","CRISTÁLIA","Leobino (Eng. Civil Abelv)",
     "Confirmação do Cliente da Aplicação de Tinta Asfáltica nos Diques do Piso Técnico, conforme orientado no Projeto.",
     None,None,date(2026,5,20),
     "20/05: Sr. Edson da Cristália na Reunião de Planejamento no Canteiro, orientou continuarmos segundo o Projeto.",
     date(2026,5,20)),

    ("PLA", date(2026,5,20),"CISTO","ABELV","Ilson (Planejador Abelv) / Victor (Planejador Abelv)",
     "Elaboração do Estudo de HH aplicado até a data 19/05/26, conforme solicitado em reunião.",
     None,None,date(2026,5,25),None,None),

    ("SMS", date(2026,5,20),"CISTO","KROOM","Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Compra ou Projeção de Compra do Pluviômetro, com objetivo de medir a quantidade de chuva (precipitação) na Obra, e registrar os impactos das atividades.",
     "Impacto: Registro de Comprovação de Impactos de Chuvas a nível de RDO.",
     None,date(2026,5,20),None,None),

    ("TUB", date(2026,5,21),"CISTO","ABELV","Celso (Gestor de Projetos Abelv)",
     "Emissão de Solicitação de Compra (SC) para Fabricação da Linha do Dreno do TRAP",
     "Impacto: Fabricação da Linha do Trap (Início Previsto: 25/05 - Atividade do Cronograma: ID 169).",
     169,date(2026,5,21),None,None),

    ("CIV", date(2026,5,21),"CISTO","ABELV","Leobino (Eng. Civil Abelv)",
     "Elaboração de CT (Consulta Técnica) a respeito da Aplicação da Tinta Asfáltica que segundo o Projeto está orientando aplicar nos Diques do Pavimento Técnico.",
     "Impacto: Pintura Asfáltica nos Diques do Primeiro Pavimento (Início Previsto: 02/06 - Atividade do Cronograma: ID 279).",
     279,date(2026,5,21),None,None),

    ("CIV", date(2026,5,21),"CISTO","ABELV","Leobino (Eng. Civil Abelv)",
     "Solicitação de Visita Técnica e envio de documentação dos Colaboradores da empresa GEOCOM (Empresa de Ensaio de Esclerometria).",
     "Impacto: Atraso do início da confecção da base dos reatores, atividade que estava prevista para o "
     "dia 21/05/26 (ID: 22), em virtude do Projeto Fornecido pela Cristália, a base da Plataforma do "
     "Reator está localizada no Local do Ralo.",
     22,date(2026,5,21),None,None),

    ("SMS", date(2026,5,21),"CISTO","ABELV","Celso (Gestor de Projetos Abelv) / Leobino (Eng. Civil Abelv)",
     "Confecção de Andaime de Proteção contra quedas de Materiais para Atividade de Furação da Laje no Prédio do Citostático, solicitado pela Krrom.",
     "Impacto: Atraso no Início da Furação da Laje do Citostático (Início Previsto: 15/06 - Atividade do Cronograma: ID 284).",
     284,date(2026,6,8),None,None),

    ("CIV", date(2026,5,21),"CISTO","ABELV","Leobino (Eng. Civil Abelv)",
     "Elaboração de CT para Definição do Projeto dos Diques Técnicos, pois o Projeto Civil orienta a "
     "construção de 2 Diques, e o Projeto Mecânico orienta apenas 1 Dique para todos os equipamentos.",
     "Impacto: Retrabalho nos Diques de Contenção do Piso Técnico, pois já foi iniciado essa atividade. "
     "(Início Previsto: 21/05 - Atividade do Cronograma: ID 276).",
     276,date(2026,5,21),None,None),

    ("SUP", date(2026,5,21),"CISTO","ABELV","Yohanna (Engenheira Mecânica Abelv)",
     "Levantamento Inicial do Quantitativo de Tintas para Obra.",
     None,None,date(2026,5,29),None,None),

    ("CIV", date(2026,5,21),"CISTO","CRISTÁLIA","Leobino (Eng. Civil Abelv)",
     "Solicitar a Cristália a Definição da Tinta Externa que será Aplicada nos Diques de Contenção do Piso Técnico",
     None,None,date(2026,5,22),None,None),
]

# ── XLSX ────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Restrições"

headers = ["ITEM","PROJETO","DISCIPLINA","ELABORAÇÃO","EMPRESA RESP.","RESPONSÁVEL",
           "DESCRIÇÃO","IMPACTO","ID IMPACT","NECESSIDADE","OBSERVAÇÃO",
           "CONCLUSÃO REAL","DIAS ATRASADO","STATUS"]
N = 14
C_DISC=3; C_ELAB=4; C_ID=9; C_NEC=10; C_OBS=11; C_CONC=12; C_DIAS=13; C_STA=14

col_widths = [5,10,12,13,14,35,50,50,10,13,35,14,10,22]
for i,w in enumerate(col_widths,1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 36

for col,h in enumerate(headers,1):
    c = ws.cell(row=1,column=col,value=h)
    c.fill=pf(C_HDR_BG); c.font=ft(bold=True,color=C_HDR_FG)
    c.alignment=al(wrap=True); c.border=bt()

disc_colors = {
    "CIV":"D6E4F0","TUB":"D5F5E3","MEC":"FAD7A0","HVA":"E8DAEF",
    "MET":"FDFEFE","ELE":"FEF9E7","SUP":"EAFAF1","GER":"F9EBEA",
    "JUR":"FDEDEC","MED":"EBF5FB","SMS":"FFF9C4","PLA":"E8F8F5",
}

for idx,row in enumerate(dados,1):
    r = idx+1
    disc,elab,proj,emp,resp,desc,imp,id_imp,nec,obs,conc = row
    st = status(nec,conc)
    da = dias(nec,conc)
    proj_bg = C_CI_BG if proj=="CISTO" else C_FA_BG
    disc_bg = disc_colors.get(disc,"FFFFFF")

    values = [idx,proj,disc,elab,emp,resp,desc,imp,id_imp,nec,obs,conc,da,st]
    for col,val in enumerate(values,1):
        c = ws.cell(row=r,column=col,value=val)
        c.border=bt(); c.alignment=al(h="left",v="center",wrap=True)

        if col in (C_ELAB,C_NEC,C_CONC) and isinstance(val,date):
            c.number_format="DD/MM/YYYY"; c.alignment=al(h="center",v="center")
        if col in (1,C_ID,C_DIAS) and val!="—":
            c.alignment=al(h="center",v="center")

        if col==C_DIAS:
            if isinstance(val,int) and val>0: c.font=ft(bold=True,color=C_AT_FG)
            elif val=="—": c.font=ft(color="808080")
            else: c.font=ft(color="375623")

        if col==C_STA:
            m={"NO PRAZO":(C_NP_BG,C_NP_FG),"ALERTA":(C_AL_BG,C_AL_FG),
               "ATRASADO":(C_AT_BG,C_AT_FG),"CONCLUÍDO":(C_CO_BG,C_CO_FG),
               "CONCLUÍDO COM ATRASO":(C_CA_BG,C_CA_FG),"—":("F2F2F2","808080")}
            bg,fg = m.get(st,("F2F2F2","808080"))
            c.fill=pf(bg); c.font=ft(bold=True,color=fg)
        elif col==C_DISC:
            c.fill=pf(disc_bg); c.font=ft(bold=True,color="1A1A1A"); c.alignment=al(h="center",v="center")
        else:
            c.fill=pf("FFE0E0") if st=="ATRASADO" else pf(proj_bg)

    ws.row_dimensions[r].height = 30

ws.freeze_panes="A2"
ws.auto_filter.ref=f"A1:{get_column_letter(N)}{len(dados)+1}"

leg_col=N+2
ws.cell(row=1,column=leg_col,value="LEGENDA").fill=pf(C_HDR_BG)
ws.cell(row=1,column=leg_col).font=ft(bold=True,color="FFFFFF")
ws.cell(row=1,column=leg_col).alignment=al()
ws.column_dimensions[get_column_letter(leg_col)].width=24
for i,(lbl,bg,fg) in enumerate([
    ("NO PRAZO",C_NP_BG,C_NP_FG),("ALERTA (vence hoje)",C_AL_BG,C_AL_FG),
    ("ATRASADO",C_AT_BG,C_AT_FG),("CONCLUÍDO",C_CO_BG,C_CO_FG),
    ("CONCLUÍDO COM ATRASO",C_CA_BG,C_CA_FG),("SEM PRAZO","F2F2F2","808080")],2):
    c=ws.cell(row=i,column=leg_col,value=lbl)
    c.fill=pf(bg); c.font=ft(bold=True,color=fg); c.alignment=al(); c.border=bt()

XLSX_OUT = "/home/user/meu-planejamento/2026-05-21_CONTROLE_CRITICO_PLA_Restricoes-Consolidadas-25098-25103_Abelv.xlsx"
wb.save(XLSX_OUT)
print(f"✅ XLSX: {XLSX_OUT}")

# ── MD ──────────────────────────────────────────────────────────────────
sc = Counter(status(d[8],d[10]) for d in dados)
proj_c = Counter(d[2] for d in dados)
disc_c = Counter(d[0] for d in dados)

disc_label = {
    "CIV":"CIVIL","TUB":"TUBULAÇÃO","MEC":"MECÂNICA","HVA":"HVAC",
    "MET":"ESTRUTURA METÁLICA","ELE":"ELÉTRICA","SUP":"SUPRIMENTOS",
    "GER":"GERENCIAL","JUR":"JURÍDICO/CONTRATO","MED":"MEDIÇÃO",
    "SMS":"SMS","PLA":"PLANEJAMENTO",
}

lines = []
lines.append("---")
lines.append("tipo: controle")
lines.append('titulo: "Restrições Consolidadas — Cristália 25098 + 25103"')
lines.append("data: 2026-05-21")
lines.append("total: 37")
lines.append('projetos: ["25098 Citostático", "25103 Farmoquímico"]')
lines.append("tags: [restricoes, controle, planejamento, 25098, 25103]")
lines.append("---")
lines.append("")
lines.append("# Restrições Consolidadas — Cristália 25098 + 25103")
lines.append("**Data de Referência:** 21/05/2026 | **Total:** 37 restrições")
lines.append("")
lines.append("---")
lines.append("")
lines.append("## Resumo de Status")
lines.append("")
lines.append("| Status | Qtde |")
lines.append("|---|:---:|")
for s in ["ATRASADO","ALERTA","CONCLUÍDO COM ATRASO","CONCLUÍDO","NO PRAZO","—"]:
    if sc.get(s,0)>0:
        lines.append(f"| {s} | {sc[s]} |")
lines.append(f"| **TOTAL** | **{len(dados)}** |")
lines.append("")
lines.append("## Por Projeto")
lines.append("")
lines.append("| Projeto | Qtde |")
lines.append("|---|:---:|")
for p,n in sorted(proj_c.items()):
    lines.append(f"| {p} | {n} |")
lines.append("")
lines.append("## Por Disciplina")
lines.append("")
lines.append("| Código | Disciplina | Qtde |")
lines.append("|---|---|:---:|")
for d,n in sorted(disc_c.items(), key=lambda x:-x[1]):
    lines.append(f"| {d} | {disc_label.get(d,d)} | {n} |")
lines.append("")
lines.append("---")
lines.append("")
lines.append("## Restrições Detalhadas")
lines.append("")

for idx,row in enumerate(dados,1):
    disc,elab,proj,emp,resp,desc,imp,id_imp,nec,obs,conc = row
    st = status(nec,conc)
    da = dias(nec,conc)
    lines.append(f"### [{idx:02d}] {desc[:70]}{'...' if len(desc)>70 else ''}")
    lines.append("")
    lines.append(f"- **Projeto:** {proj} | **Disciplina:** {disc} — {disc_label.get(disc,disc)}")
    lines.append(f"- **Empresa Resp.:** {emp} | **Responsável:** {resp}")
    lines.append(f"- **Elaboração:** {elab.strftime('%d/%m/%Y')}")
    if nec: lines.append(f"- **Necessidade:** {nec.strftime('%d/%m/%Y')}")
    if conc: lines.append(f"- **Conclusão Real:** {conc.strftime('%d/%m/%Y')}")
    if id_imp: lines.append(f"- **ID Impact (Cronograma):** {id_imp}")
    lines.append(f"- **Status:** {st}" + (f" ({da} dias)" if isinstance(da,int) and da>0 else ""))
    if imp: lines.append(f"- **Impacto:** {imp}")
    if obs: lines.append(f"- **Observação:** {obs}")
    lines.append("")

lines.append("---")
lines.append("")
lines.append("*Gerado automaticamente em 21/05/2026 | Fonte: Planilha original do usuário — apenas correções de Português e Pontuação aplicadas.*")

MD_OUT = "/home/user/meu-planejamento/2026-05-21_CONTROLE_CRITICO_PLA_Restricoes-Consolidadas-25098-25103_Abelv.md"
with open(MD_OUT,"w",encoding="utf-8") as f:
    f.write("\n".join(lines))
print(f"✅ MD:   {MD_OUT}")
