#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Planilha Consolidada de Restrições — Projetos Cristália (Controle Interno)
25098 Citostático (CISTO) + 25103 Farmoquímico (FARMO)
Data de referência: 20/05/2026
Fonte: Planilha original do usuário — apenas correções de Português e Pontuação aplicadas.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
from collections import Counter

# ============================================================
# HELPERS
# ============================================================
def pf(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def ft(bold=False, color="000000", size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Calibri")

def al(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border_thin():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

# ============================================================
# CORES
# ============================================================
C_HEADER_BG   = "1F4E79"
C_HEADER_FG   = "FFFFFF"
C_NO_PRAZO_BG = "C6EFCE"
C_NO_PRAZO_FG = "375623"
C_ALERTA_BG   = "FFEB9C"
C_ALERTA_FG   = "7F6000"
C_ATRASADO_BG = "FFC7CE"
C_ATRASADO_FG = "9C0006"
C_CONC_BG     = "DAEEF3"
C_CONC_FG     = "17375E"
C_CONC_ATR_BG = "FFF2CC"
C_CONC_ATR_FG = "7F4F00"
C_CISTO_BG    = "EBF3FB"
C_FARMO_BG    = "FFF2CC"

# ============================================================
# CÁLCULOS
# ============================================================
today = date(2026, 5, 20)

def calc_status(necessidade, conclusao):
    if necessidade is None:
        return "—"
    if conclusao is not None:
        return "CONCLUÍDO" if conclusao <= necessidade else "CONCLUÍDO COM ATRASO"
    if necessidade < today:
        return "ATRASADO"
    if necessidade == today:
        return "ALERTA"
    return "NO PRAZO"

def calc_dias(necessidade, conclusao):
    if necessidade is None:
        return "—"
    if conclusao is not None:
        return max(0, (conclusao - necessidade).days)
    if necessidade < today:
        return (today - necessidade).days
    return 0

# ============================================================
# DADOS — 24 REGISTROS
# Fonte: Planilha original do usuário (c3fe8ee3)
# Apenas correções de Português e Pontuação aplicadas.
# Tupla: (elaboracao, projeto, empresa_resp, responsavel,
#         descricao, impacto, necessidade, observacao, conclusao_real)
# ============================================================
dados = [
    # 1 — CISTO / KROOM
    (date(2026,5,11), "CISTO", "KROOM",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Mobilização da Krrom para realização da Tarefa do Dique no Piso Técnico",
     None,
     date(2026,5,20), None, date(2026,5,11)),

    # 2 — CISTO / KROOM
    (date(2026,5,8), "CISTO", "KROOM",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Recebimento de Materiais do Dique no Piso Técnico",
     None,
     date(2026,5,11), None, date(2026,5,11)),

    # 3 — CISTO / CRISTÁLIA  [corr: imepdindo→impedindo, continidade→continuidade, da Bases→das Bases, Impedi→Impede]
    (date(2026,5,8), "CISTO", "CRISTÁLIA",
     "Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Retirada de 2 Peças de Vidro da Sala Limpa, pois está impedindo a continuidade da Marcação das Bases que serão construídas.",
     "Impede a continuação da Locação das Bases que serão construídas na Sala Limpa.",
     date(2026,5,19), None, None),

    # 4 — CISTO / CRISTÁLIA  [corr: do Reatores→dos Reatores, da das bases→das Bases]
    (date(2026,5,14), "CISTO", "CRISTÁLIA",
     "Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Cobrar Definição da Cristália a respeito da Locação da Plataforma dos Reatores em Função da Instalação da Nova Linha Hidrossanitária.",
     "Impacta o início do Corte do Piso para confecção das Bases da Plataforma dos Reatores "
     "(Início Previsto: 21/05- Atividade do Cronograma: ID 22). E Demolição do Piso para "
     "Instalação da Rede Hidrossanitária (Início Previsto: 18/05 - Atividade do Cronograma: ID 58).",
     date(2026,5,18), None, None),

    # 5 — CISTO / ABELV
    (date(2026,5,12), "CISTO", "ABELV",
     "Tiago (Gestor de Projetos Abelv)",
     "Acompanhamento / Recebimento das Conexões de Aço Carbono",
     "Impacta na Fabricação da Tubulação de Aço Carbono "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 54 e ID 317).",
     date(2026,5,19), None, None),

    # 6 — CISTO / ABELV
    (date(2026,5,12), "CISTO", "ABELV",
     "Tiago (Gestor de Projetos Abelv)",
     "Acompanhamento / Recebimento da Tubulação de Aço Carbono",
     "Impacta na Fabricação da Tubulação de Aço Carbono "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 54 e ID 317).",
     date(2026,5,15), None, date(2026,5,15)),

    # 7 — CISTO / ABELV  [corr: Lavadaor→Lavador, Inicio→Início, da Piso→do Piso]
    (date(2026,5,15), "CISTO", "ABELV",
     "Uehara (Engenheiro Mecânico do HVAC Abelv)",
     "Validar a Mudança Dimensional da Base do Lavador de Gases",
     "Impacta no Início da Demolição do Piso do Radier do Lavador de Gases a Título de Área "
     "caso aconteça alguma mudança significativa. "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 466).",
     date(2026,5,20), None, None),

    # 8 — CISTO / ABELV
    (date(2026,5,12), "CISTO", "ABELV",
     "Uehara (Engenheiro Mecânico do HVAC Abelv)",
     "Contratação Lavador de Gases",
     None,
     date(2026,5,15),
     "19/05: Aguardando o Setor de Suprimentos da Abelv Fechar o Pedido.",
     None),

    # 9 — CISTO / ABELV (Ana Assis)  [corr: Lavadaor→Lavador, imapctando→impactando, Inicio→Início, da Piso→do Piso]
    (date(2026,5,19), "CISTO", "ABELV",
     "Ana Assis (Gerente do Suprimentos Abelv)",
     "Fechamento do Pedido de Compra do Lavador de Gases",
     "Impacta na identificação real do Dimensional da Base do Lavador de Gases, "
     "consequentemente também impactando o Início da Demolição do Piso do Radier do Lavador de Gases "
     "a Título de Área caso aconteça alguma mudança significativa. "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 466).",
     date(2026,5,19), None, None),

    # 10 — FARMO / KROOM  [corr: espaço duplo em "Tipo de  Fundação"]
    (date(2026,5,12), "FARMO", "KROOM",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Definição do Tipo de Fundação do Prédio",
     "Impacta na Contratação da Empresa para realização da Estaca, por ventura impactando na compra "
     "de materiais da estaca (Início Previsto: 08/06 - Atividade do Cronograma: ID 118).",
     date(2026,5,13), None, None),

    # 11 — FARMO / CRISTÁLIA  [corr: Inicio→Início]
    (date(2026,5,12), "FARMO", "CRISTÁLIA",
     "Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Recebimento do Projeto da Base do Lavador de Gases Existente.",
     "Impacta na Formalização da Nova Locação que foi combinada IN LOCO para execução da base do "
     "Lavador existente. Início da Confecção da Nova Base (Radier) do Lavador de Gases Existente "
     "(Início Previsto: 05/06/26 - Atividade do Cronograma: ID 251).",
     date(2026,5,15), None, None),

    # 12 — FARMO / MGC (Yohanna)  [corr: Engenheiro Mecânica→Engenheira Mecânica, Subscontratações→Subcontratações]
    (date(2026,5,7), "FARMO", "MGC",
     "Yohanna (Engenheira Mecânica Abelv)",
     "Solicitação e Recebimento dos Projetos do Farmoquímico da Empresa MGC, e envio para Cristália para aprovação",
     "Impacta na Compra dos Materiais e Subcontratações para realização dos materiais.",
     date(2026,5,19), None, None),

    # 13 — FARMO / ABELV (Leobino cotas Lavador)  [corr: Lovador→Lavador]
    (date(2026,5,19), "FARMO", "ABELV",
     "Leobino (Eng. Civil Abelv)",
     "Enviar para a MGC as cotas de elevação do Lavador de Gases",
     "Impacta na confecção do Projeto do Radier do Lavador de Gases",
     date(2026,5,19), None, None),

    # 14 — FARMO / ABELV (Leobino cotas Sala Elétrica)
    (date(2026,5,19), "FARMO", "ABELV",
     "Leobino (Eng. Civil Abelv)",
     "Enviar para a MGC as cotas de elevação da Sala Elétrica",
     "Impacta na confecção do Projeto da Sala Elétrica",
     date(2026,5,20), None, None),

    # 15 — FARMO / KROOM (data fundação)  [corr: informa→informar, do que→em que]
    (date(2026,5,19), "FARMO", "KROOM",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "A Krrom precisa informar a Data em que irão apresentar o Tipo de Fundação que será aplicada no Farmoquímico.",
     "Impacta na Contratação da Empresa para realização da Estaca, por ventura impactando na compra "
     "de materiais da estaca (Início Previsto: 08/06 - Atividade do Cronograma: ID 118).",
     date(2026,5,19), None, None),

    # 16 — CISTO / KROOM (Metodologia de Corte)  [corr: resistencia→resistência, contruido→construída, da das→das, do Reatores→dos Reatores]
    (date(2026,5,19), "CISTO", "KROOM",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "A Krrom precisa apresentar a Metodologia de Corte que será aplicado para avaliar a resistência "
     "do piso onde poderá ser construída a base da Plataforma dos Reatores, em função da Interferência "
     "da Instalação da Nova Rede Hidrossanitária.",
     "Impacta o início do Corte do Piso para confecção das Bases da Plataforma dos Reatores "
     "(Início Previsto: 21/05- Atividade do Cronograma: ID 22). E Demolição do Piso para "
     "Instalação da Rede Hidrossanitária (Início Previsto: 18/05 - Atividade do Cronograma: ID 58).",
     date(2026,5,20), None, None),

    # 17 — FARMO / KROOM (Tubulação Pluvial)
    (date(2026,5,19), "FARMO", "KROOM",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "Elaboração do Pedido da Tubulação de Água Pluvial",
     "Impacta a Compra da Tubulação de Água Pluvial",
     date(2026,5,27), None, None),

    # 18 — FARMO / KROOM (Rodrigo - Faturamento)  [corr: Imapcto→Impacto]
    (date(2026,5,19), "FARMO", "KROOM",
     "Rodrigo (Téc. Planej. da Kroom) / Leobino (Eng. Civil Abelv)",
     "A Kroom precisa enviar os Pedidos de Faturamento Direto",
     "Impacto na Medição",
     date(2026,5,19), None, None),

    # 19 — CISTO / KROOM (Andaime)
    (date(2026,5,19), "CISTO", "KROOM",
     "Tiago (Gestor de Projetos Abelv) / Celso (Gestor de Projetos Abelv)",
     "Verificação do Efetivo de Andaime para atendimento das Obras (Farmoquímico e Citostático).",
     "Impacto de acesso para realização das Frentes de Montagem",
     date(2026,5,20), None, None),

    # 20 — CISTO / ABELV (Pipe-Shop)
    (date(2026,5,19), "CISTO", "ABELV",
     "Leobino (Eng. Civil Abelv)",
     "Montagem do Pipe-Shop de Aço Carbono no Canteiro de Obra Abelv.",
     "Impacta na Fabricação da Tubulação de Aço Carbono "
     "(Início Previsto: 25/05/26 - Atividade do Cronograma: ID 54 e ID 317).",
     date(2026,5,20), None, None),

    # 21 — CISTO / ABELV (Erik Massola - Açoplast)  [corr: EriK→Erik, Tiago Tiago→Tiago]
    (date(2026,5,19), "CISTO", "ABELV",
     "Erik Massola (Comprador Abelv) / Tiago (Gestor de Projetos Abelv)",
     "O Setor de Compras precisa enviar o Efetivo (Nomes e Cpfs dos colaboradores) da empresa "
     "contratada Açoplast (Estrutura Metálica) para Tiago Matos, para posterior envio para Cristália "
     "realizar a avaliação social dos mesmos.",
     "Impacto: Caso os mesmos não forem aprovados, eles não poderão adentrar ao site, podendo "
     "impactar o prazo da obra.",
     date(2026,5,20), None, None),

    # 22 — CISTO / ABELV (Edimar Cunha - Andaime)  [corr: Tiago Tiago→Tiago]
    (date(2026,5,19), "CISTO", "ABELV",
     "Edimar Cunha (Supervisor de Suprimentos Abelv) / Tiago (Gestor de Projetos Abelv)",
     "Contratação do Fornecedor de Andaime.",
     "Impacto de acesso para realização das Frentes de Montagem",
     date(2026,5,22), None, None),

    # 23 — FARMO / ABELV (Recicla - Documentação)  [corr: aditoria→auditoria]
    (date(2026,5,19), "FARMO", "ABELV",
     "Diego (Gerente de Projetos da Kroom) / Leobino (Eng. Civil Abelv)",
     "A empresa Krrom precisa Regularizar a Documentação da Empresa Recicla (Bota Fora).",
     "Impacto: Caso aconteça alguma auditoria ou acidente, estamos descobertos por documentação "
     "legal de subcontratação.",
     date(2026,5,21), None, None),

    # 24 — CISTO / ABELV (Yohanna - Plataforma)  [corr: Engenheiro→Engenheira, não cabe→não cabem]
    (date(2026,5,19), "CISTO", "ABELV",
     "Yohanna (Engenheira Mecânica Abelv)",
     "Verificação da readequação do Projeto da Plataforma dos Reatores por base do nosso fornecedor "
     "que vai fabricar, pois os reatores adquiridos pela Cristália não cabem no mesmo.",
     None,
     None, None, None),
]

# ============================================================
# WORKBOOK
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Restrições"

headers = [
    "ITEM", "PROJETO", "ELABORAÇÃO", "EMPRESA RESP.", "RESPONSÁVEL",
    "DESCRIÇÃO", "IMPACTO", "NECESSIDADE", "OBSERVAÇÃO",
    "CONCLUSÃO REAL", "DIAS ATRASADO", "STATUS"
]
N_COLS     = 12
COL_STATUS = 12
COL_DIAS   = 11
COL_NEC    = 8
COL_CONC   = 10

col_widths = [5, 10, 13, 14, 35, 50, 50, 13, 35, 14, 10, 22]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 36

for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.fill      = pf(C_HEADER_BG)
    c.font      = ft(bold=True, color=C_HEADER_FG, size=10)
    c.alignment = al(h="center", v="center", wrap=True)
    c.border    = border_thin()

for idx, row in enumerate(dados, 1):
    r = idx + 1
    elab, proj, emp, resp, desc, imp, nec, obs, conc = row

    status   = calc_status(nec, conc)
    dias_atr = calc_dias(nec, conc)
    proj_bg  = C_CISTO_BG if proj == "CISTO" else C_FARMO_BG

    values = [idx, proj, elab, emp, resp, desc, imp, nec, obs, conc, dias_atr, status]

    for col, val in enumerate(values, 1):
        c = ws.cell(row=r, column=col, value=val)
        c.border    = border_thin()
        c.alignment = al(h="left", v="center", wrap=True)

        if col in (COL_NEC, COL_CONC) and isinstance(val, date):
            c.number_format = "DD/MM/YYYY"
            c.alignment = al(h="center", v="center")

        if col in (1, COL_DIAS) and val != "—":
            c.alignment = al(h="center", v="center")

        if col == COL_DIAS:
            if isinstance(val, int) and val > 0:
                c.font = ft(bold=True, color=C_ATRASADO_FG)
            elif val == "—":
                c.font = ft(color="808080")
            else:
                c.font = ft(color="375623")

        if col == COL_STATUS:
            if status == "NO PRAZO":
                c.fill = pf(C_NO_PRAZO_BG); c.font = ft(bold=True, color=C_NO_PRAZO_FG)
            elif status == "ALERTA":
                c.fill = pf(C_ALERTA_BG);   c.font = ft(bold=True, color=C_ALERTA_FG)
            elif status == "ATRASADO":
                c.fill = pf(C_ATRASADO_BG); c.font = ft(bold=True, color=C_ATRASADO_FG)
            elif status == "CONCLUÍDO":
                c.fill = pf(C_CONC_BG);     c.font = ft(bold=True, color=C_CONC_FG)
            elif status == "CONCLUÍDO COM ATRASO":
                c.fill = pf(C_CONC_ATR_BG); c.font = ft(bold=True, color=C_CONC_ATR_FG)
            else:
                c.fill = pf("F2F2F2");       c.font = ft(color="808080", bold=True)
        else:
            if status == "ATRASADO":
                c.fill = pf("FFE0E0")
            else:
                c.fill = pf(proj_bg)

    ws.row_dimensions[r].height = 30

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(N_COLS)}{len(dados)+1}"

# Legenda
legenda_col = N_COLS + 2
ws.cell(row=1, column=legenda_col, value="LEGENDA DE STATUS").fill = pf(C_HEADER_BG)
ws.cell(row=1, column=legenda_col).font = ft(bold=True, color="FFFFFF")
ws.cell(row=1, column=legenda_col).alignment = al()
ws.column_dimensions[get_column_letter(legenda_col)].width = 24

legendas = [
    ("NO PRAZO",             C_NO_PRAZO_BG, C_NO_PRAZO_FG),
    ("ALERTA (vence hoje)",  C_ALERTA_BG,   C_ALERTA_FG),
    ("ATRASADO",             C_ATRASADO_BG, C_ATRASADO_FG),
    ("CONCLUÍDO",            C_CONC_BG,     C_CONC_FG),
    ("CONCLUÍDO COM ATRASO", C_CONC_ATR_BG, C_CONC_ATR_FG),
    ("SEM PRAZO DEFINIDO",   "F2F2F2",      "808080"),
]
for i, (lbl, bg, fg) in enumerate(legendas, 2):
    c = ws.cell(row=i, column=legenda_col, value=lbl)
    c.fill = pf(bg); c.font = ft(bold=True, color=fg)
    c.alignment = al(h="center"); c.border = border_thin()

# ============================================================
# DASHBOARD
# ============================================================
wd = wb.create_sheet("Dashboard")
wd.sheet_view.showGridLines = False

def dash_cell(ws, row, col, val, bg="FFFFFF", bold=False, color="000000", h="left"):
    c = ws.cell(row=row, column=col, value=val)
    c.fill = pf(bg); c.font = ft(bold=bold, color=color)
    c.alignment = al(h=h); c.border = border_thin()
    return c

for i, w in enumerate([3, 30, 14, 3, 22, 10, 3, 22, 10], 1):
    wd.column_dimensions[get_column_letter(i)].width = w

wd.merge_cells("B1:I1")
c = wd.cell(row=1, column=2, value="DASHBOARD — RESTRIÇÕES | Cristália 25098 + 25103 | 20/05/2026")
c.fill = pf(C_HEADER_BG); c.font = ft(bold=True, color="FFFFFF", size=13)
c.alignment = al(h="center"); wd.row_dimensions[1].height = 30

status_list = [calc_status(d[6], d[8]) for d in dados]
status_count = Counter(status_list)
order_status = ["ATRASADO", "ALERTA", "CONCLUÍDO COM ATRASO", "CONCLUÍDO", "NO PRAZO", "—"]
status_colors = {
    "ATRASADO":             (C_ATRASADO_BG, C_ATRASADO_FG),
    "ALERTA":               (C_ALERTA_BG,   C_ALERTA_FG),
    "CONCLUÍDO COM ATRASO": (C_CONC_ATR_BG, C_CONC_ATR_FG),
    "CONCLUÍDO":            (C_CONC_BG,     C_CONC_FG),
    "NO PRAZO":             (C_NO_PRAZO_BG, C_NO_PRAZO_FG),
    "—":                    ("F2F2F2",       "808080"),
}

wd.merge_cells("B3:C3")
dash_cell(wd, 3, 2, "POR STATUS — QTDE", bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")
for i, s in enumerate(order_status, 4):
    bg, fg = status_colors[s]
    dash_cell(wd, i, 2, s,                     bg=bg, bold=True, color=fg)
    dash_cell(wd, i, 3, status_count.get(s, 0), bg=bg, bold=True, color=fg, h="center")

total_row = 4 + len(order_status)
dash_cell(wd, total_row, 2, "TOTAL",     bg="1F4E79", bold=True, color="FFFFFF")
dash_cell(wd, total_row, 3, len(dados),  bg="1F4E79", bold=True, color="FFFFFF", h="center")

proj_count = Counter(d[1] for d in dados)
wd.merge_cells("E3:F3")
dash_cell(wd, 3, 5, "POR PROJETO — QTDE", bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")
proj_colors = {"CISTO": (C_CISTO_BG, "17375E"), "FARMO": (C_FARMO_BG, "7F4F00")}
for i, (proj, cnt) in enumerate(sorted(proj_count.items()), 4):
    bg, fg = proj_colors.get(proj, ("FFFFFF", "000000"))
    dash_cell(wd, i, 5, proj, bg=bg, bold=True, color=fg)
    dash_cell(wd, i, 6, cnt,  bg=bg, bold=True, color=fg, h="center")

resp_atr = Counter(d[3] for d in dados if calc_status(d[6], d[8]) == "ATRASADO")
wd.row_dimensions[12].height = 8
wd.merge_cells("B13:C13")
dash_cell(wd, 13, 2, "RESPONSÁVEIS COM MAIS ITENS ATRASADOS",
          bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")
for i, (resp, cnt) in enumerate(resp_atr.most_common(8), 14):
    dash_cell(wd, i, 2, resp, bg=C_ATRASADO_BG, color=C_ATRASADO_FG)
    dash_cell(wd, i, 3, cnt,  bg=C_ATRASADO_BG, bold=True, color=C_ATRASADO_FG, h="center")

alertas = [(d[3], d[4], d[6]) for d in dados if calc_status(d[6], d[8]) == "ALERTA"]
wd.merge_cells("E13:F13")
dash_cell(wd, 13, 5, "⚠ ALERTA — VENCE HOJE (20/05)",
          bg=C_ALERTA_BG, bold=True, color=C_ALERTA_FG, h="center")
for i, (resp, desc, nec) in enumerate(alertas, 14):
    dash_cell(wd, i, 5, resp, bg=C_ALERTA_BG, color=C_ALERTA_FG)
    txt = desc[:60] + "…" if len(desc) > 60 else desc
    c = wd.cell(row=i, column=6, value=txt)
    c.fill = pf(C_ALERTA_BG); c.font = ft(color=C_ALERTA_FG)
    c.border = border_thin(); c.alignment = al(h="left", wrap=True)
    wd.row_dimensions[i].height = 28

# ============================================================
# SALVAR
# ============================================================
output_path = "/home/user/meu-planejamento/Restricoes_Consolidadas_Cristalia_20052026.xlsx"
wb.save(output_path)
print(f"✅ Arquivo salvo: {output_path}")
print(f"   Total de registros: {len(dados)}")
status_summary = dict(status_count)
print(f"   Status: {status_summary}")
