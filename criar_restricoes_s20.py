#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Planilha de Restrições — Reunião de Planejamento S20
Projetos Cristália — 25098 Citostático + 25103 Farmoquímico
Versão modificada para apresentação ao cliente
Gerado por: Claude Conector — Ministro das Pontes Digitais
Data: 20/05/2026
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
from collections import Counter

# ============================================================
# HELPERS
# ============================================================
def pf(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def ft(bold=False, color="000000", size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Calibri")

def al(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border_thin():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

# ============================================================
# CORES
# ============================================================
C_HEADER_BG   = "1F4E79"
C_HEADER_FG   = "FFFFFF"
C_NO_PRAZO_BG = "C6EFCE"
C_NO_PRAZO_FG = "375623"
C_ALERTA_BG   = "FFEB9C"
C_ALERTA_FG   = "7F6000"
C_ATRASADO_BG = "FFC7CE"
C_ATRASADO_FG = "9C0006"
C_CISTO_BG    = "EBF3FB"
C_FARMO_BG    = "FFF2CC"

# ============================================================
# DATA E CÁLCULOS
# ============================================================
today = date(2026, 5, 20)

def calc_status(necessidade, conclusao):
    if necessidade is None:
        return "—"
    if conclusao is not None:
        return "CONCLUÍDO" if conclusao <= necessidade else "CONCLUÍDO COM ATRASO"
    if necessidade < today:
        return "ATRASADO"
    if necessidade == today:
        return "ALERTA"
    return "NO PRAZO"

def calc_dias(necessidade, conclusao):
    if necessidade is None:
        return "—"
    if conclusao is not None:
        return max(0, (conclusao - necessidade).days)
    if necessidade < today:
        return (today - necessidade).days
    return 0

# ============================================================
# DADOS — 9 RESTRIÇÕES
# ============================================================
# (elaboracao, projeto, empresa_resp, responsavel,
#  descricao, impacto, necessidade, observacao, conclusao_real)

dados = [
    # Item 1
    (date(2026, 5, 8), "CISTO", "CRISTÁLIA",
     "Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Retirada de 2 Peças de Vidro da Sala Limpa, pois está impedindo a continuidade da Marcação das Bases que serão construídas.",
     "Impede a continuação da Locação das Bases que serão construídas na Sala Limpa.",
     date(2026, 5, 19), None, None),

    # Item 2
    (date(2026, 5, 14), "CISTO", "CRISTÁLIA",
     "Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Cobrar Definição da Cristália a respeito da Locação da Plataforma dos Reatores em Função da Instalação da Nova Linha Hidrossanitária.",
     "Impacta o início do Corte do Piso para confecção das Bases da Plataforma dos Reatores (Início Previsto: 21/05 - Atividade do Cronograma: ID 22). E Demolição do Piso para Instalação da Rede Hidrossanitária (Início Previsto: 18/05 - Atividade do Cronograma: ID 58).",
     date(2026, 5, 18), None, None),

    # Item 3
    (date(2026, 5, 15), "CISTO", "ABELV",
     "Uehara (Engenheiro Mecânico do HVAC Abelv)",
     "Validar a Mudança Dimensional da Base do Lavador de Gases.",
     "Impacta no Início da Demolição do Piso do Radier do Lavador de Gases a Título de Área, caso aconteça alguma mudança significativa (Início Previsto: 25/05/26 - Atividade do Cronograma: ID 466).",
     date(2026, 5, 20), None, None),

    # Item 4
    (date(2026, 5, 12), "FARMO", "KRROM",
     "Diego (Gerente de Projetos da Krrom) / Leobino (Eng. Civil Abelv)",
     "Definição do Tipo de Fundação do Prédio.",
     "Impacta na Contratação da Empresa para realização da Estaca, por ventura impactando na compra de materiais da estaca (Início Previsto: 08/06 - Atividade do Cronograma: ID 118).",
     date(2026, 5, 13), None, None),

    # Item 5
    (date(2026, 5, 12), "FARMO", "CRISTÁLIA",
     "Edson (Cristália) / Leobino (Eng. Civil Abelv)",
     "Recebimento do Projeto da Base do Lavador de Gases Existente.",
     "Impacta na Formalização da Nova Locação que foi combinada IN LOCO para execução da base do Lavador existente. Início da Confecção da Nova Base (Radier) do Lavador de Gases Existente (Início Previsto: 05/06/26 - Atividade do Cronograma: ID 251).",
     date(2026, 5, 15), None, None),

    # Item 6
    (date(2026, 5, 7), "FARMO", "MGC",
     "Yohanna (Engenheira Mecânica Abelv)",
     "Solicitação e Recebimento dos Projetos do Farmoquímico da Empresa MGC, e envio para Cristália para aprovação.",
     "Impacta na Compra dos Materiais e Subcontratações para realização dos serviços.",
     date(2026, 5, 19), None, None),

    # Item 7
    (date(2026, 5, 19), "FARMO", "KRROM",
     "Diego (Gerente de Projetos da Krrom) / Leobino (Eng. Civil Abelv)",
     "A Krrom precisa informar a data em que irá apresentar o Tipo de Fundação que será aplicada no Farmoquímico.",
     "Impacta na Contratação da Empresa para realização da Estaca, por ventura impactando na compra de materiais da estaca (Início Previsto: 08/06 - Atividade do Cronograma: ID 118).",
     date(2026, 5, 19), None, None),

    # Item 8
    (date(2026, 5, 19), "CISTO", "KRROM",
     "Diego (Gerente de Projetos da Krrom) / Leobino (Eng. Civil Abelv)",
     "A Krrom precisa apresentar a Metodologia de Corte que será aplicada para avaliar a resistência do piso onde poderá ser construída a base da Plataforma dos Reatores, em função da Interferência da Instalação da Nova Rede Hidrossanitária.",
     "Impacta o início do Corte do Piso para confecção das Bases da Plataforma dos Reatores (Início Previsto: 21/05 - Atividade do Cronograma: ID 22). E Demolição do Piso para Instalação da Rede Hidrossanitária (Início Previsto: 18/05 - Atividade do Cronograma: ID 58).",
     date(2026, 5, 20), None, None),

    # Item 10 → Item 9 (sem prazo definido)
    (date(2026, 5, 19), "CISTO", "ABELV",
     "Yohanna (Engenheira Mecânica Abelv)",
     "Verificação da readequação do Projeto da Plataforma dos Reatores por base do nosso fornecedor que vai fabricar, pois os reatores adquiridos pela Cristália não cabem no mesmo.",
     "—",
     None, None, None),
]

# ============================================================
# WORKBOOK
# ============================================================
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Restrições"

# ── Cabeçalhos ───────────────────────────────────────────────
headers = [
    "ITEM", "PROJETO", "ELABORAÇÃO", "EMPRESA\nRESP.",
    "RESPONSÁVEL", "DESCRIÇÃO", "IMPACTO",
    "NECESSIDADE", "OBSERVAÇÃO", "CONCLUSÃO\nREAL",
    "DIAS\nATRASADO", "STATUS"
]

COL_NECESSIDADE = 8
COL_CONCLUSAO   = 10
COL_DIAS        = 11
COL_STATUS      = 12
N_COLS          = 12

col_widths = [6, 8, 13, 16, 26, 50, 55, 13, 30, 14, 10, 22]

for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 36

for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.fill      = pf(C_HEADER_BG)
    c.font      = ft(bold=True, color=C_HEADER_FG, size=10)
    c.alignment = al(h="center", v="center", wrap=True)
    c.border    = border_thin()

# ── Dados ────────────────────────────────────────────────────
for idx, row in enumerate(dados, 1):
    r = idx + 1
    elab, proj, empresa, resp, desc, imp, nec, obs, conc = row

    status   = calc_status(nec, conc)
    dias_atr = calc_dias(nec, conc)

    proj_bg = C_CISTO_BG if proj == "CISTO" else C_FARMO_BG

    values = [
        idx, proj, elab, empresa, resp,
        desc, imp, nec, obs if obs else "",
        conc, dias_atr, status
    ]

    for col, val in enumerate(values, 1):
        c = ws.cell(row=r, column=col, value=val)
        c.border    = border_thin()
        c.alignment = al(h="left", v="center", wrap=True)

        # Datas centradas
        if col in (3, COL_NECESSIDADE, COL_CONCLUSAO) and isinstance(val, date):
            c.number_format = "DD/MM/YYYY"
            c.alignment = al(h="center", v="center")

        # Números centrados
        if col in (1, COL_DIAS):
            c.alignment = al(h="center", v="center")

        # Cor dos dias de atraso
        if col == COL_DIAS:
            if isinstance(dias_atr, int) and dias_atr > 0:
                c.font = ft(bold=True, color=C_ATRASADO_FG)
            else:
                c.font = ft(color="375623")

        # Cor do status
        if col == COL_STATUS:
            if status == "NO PRAZO":
                c.fill = pf(C_NO_PRAZO_BG)
                c.font = ft(bold=True, color=C_NO_PRAZO_FG)
            elif status == "ALERTA":
                c.fill = pf(C_ALERTA_BG)
                c.font = ft(bold=True, color=C_ALERTA_FG)
            elif status == "ATRASADO":
                c.fill = pf(C_ATRASADO_BG)
                c.font = ft(bold=True, color=C_ATRASADO_FG)
            elif status in ("CONCLUÍDO", "CONCLUÍDO COM ATRASO"):
                c.fill = pf("DAEEF3")
                c.font = ft(bold=True, color="17375E")
            else:
                c.fill = pf(proj_bg)
                c.font = ft(color="404040", italic=True)
        elif col != COL_STATUS:
            if status == "ATRASADO":
                c.fill = pf("FFE0E0")
            else:
                c.fill = pf(proj_bg)

    ws.row_dimensions[r].height = 45

# ── Congelar + Auto-filtro ────────────────────────────────────
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(N_COLS)}{len(dados)+1}"

# ── Legenda ───────────────────────────────────────────────────
legenda_col = N_COLS + 2
ws.cell(row=1, column=legenda_col, value="LEGENDA DE STATUS").fill = pf(C_HEADER_BG)
ws.cell(row=1, column=legenda_col).font = ft(bold=True, color="FFFFFF")
ws.cell(row=1, column=legenda_col).alignment = al()
ws.column_dimensions[get_column_letter(legenda_col)].width = 24

legendas = [
    ("NO PRAZO",             C_NO_PRAZO_BG, C_NO_PRAZO_FG),
    ("ALERTA (vence hoje)",  C_ALERTA_BG,   C_ALERTA_FG),
    ("ATRASADO",             C_ATRASADO_BG, C_ATRASADO_FG),
    ("CONCLUÍDO",            "DAEEF3",      "17375E"),
    ("SEM PRAZO DEFINIDO",   "F2F2F2",      "404040"),
]
for i, (lbl, bg, fg) in enumerate(legendas, 2):
    c = ws.cell(row=i, column=legenda_col, value=lbl)
    c.fill      = pf(bg)
    c.font      = ft(bold=True, color=fg)
    c.alignment = al(h="center")
    c.border    = border_thin()

# ============================================================
# ABA DASHBOARD
# ============================================================
wd = wb.create_sheet("Dashboard")
wd.sheet_view.showGridLines = False

def dash_cell(ws, row, col, val, bg="FFFFFF", bold=False, color="000000", h="left"):
    c = ws.cell(row=row, column=col, value=val)
    c.fill      = pf(bg)
    c.font      = ft(bold=bold, color=color)
    c.alignment = al(h=h)
    c.border    = border_thin()
    return c

for i, w in enumerate([3, 30, 14, 3, 22, 10], 1):
    wd.column_dimensions[get_column_letter(i)].width = w

wd.merge_cells("B1:F1")
c = wd.cell(row=1, column=2,
            value="DASHBOARD — RESTRIÇÕES | Cristália 25098 + 25103 | Reunião de Planejamento S20 | 20/05/2026")
c.fill      = pf(C_HEADER_BG)
c.font      = ft(bold=True, color="FFFFFF", size=12)
c.alignment = al(h="center")
wd.row_dimensions[1].height = 30

# Por Status
status_list  = [calc_status(d[6], d[8]) for d in dados]
status_count = Counter(s for s in status_list if s != "—")

wd.merge_cells("B3:C3")
dash_cell(wd, 3, 2, "POR STATUS — QTDE", bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")

order_status = ["ATRASADO", "ALERTA", "NO PRAZO", "CONCLUÍDO"]
status_colors = {
    "ATRASADO": (C_ATRASADO_BG, C_ATRASADO_FG),
    "ALERTA":   (C_ALERTA_BG,   C_ALERTA_FG),
    "NO PRAZO": (C_NO_PRAZO_BG, C_NO_PRAZO_FG),
    "CONCLUÍDO":("DAEEF3",      "17375E"),
}
for i, s in enumerate(order_status, 4):
    bg, fg = status_colors[s]
    dash_cell(wd, i, 2, s,                    bg=bg, bold=True, color=fg)
    dash_cell(wd, i, 3, status_count.get(s,0), bg=bg, bold=True, color=fg, h="center")

sem_prazo = sum(1 for s in status_list if s == "—")
dash_cell(wd, 4+len(order_status), 2, "SEM PRAZO DEFINIDO", bg="F2F2F2", color="404040")
dash_cell(wd, 4+len(order_status), 3, sem_prazo, bg="F2F2F2", color="404040", h="center")

total_row = 5 + len(order_status)
dash_cell(wd, total_row, 2, "TOTAL",    bg=C_HEADER_BG, bold=True, color="FFFFFF")
dash_cell(wd, total_row, 3, len(dados), bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")

# Por Projeto
proj_count = Counter(d[1] for d in dados)
wd.merge_cells("E3:F3")
dash_cell(wd, 3, 5, "POR PROJETO — QTDE", bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")
proj_colors = {"CISTO": (C_CISTO_BG,"17375E"), "FARMO": (C_FARMO_BG,"7F4F00")}
for i, (proj, cnt) in enumerate(sorted(proj_count.items()), 4):
    bg, fg = proj_colors.get(proj, ("FFFFFF","000000"))
    dash_cell(wd, i, 5, proj, bg=bg, bold=True, color=fg)
    dash_cell(wd, i, 6, cnt,  bg=bg, bold=True, color=fg, h="center")

# Atrasados por empresa
wd.row_dimensions[10].height = 8
wd.merge_cells("B11:C11")
dash_cell(wd, 11, 2, "ITENS ATRASADOS POR EMPRESA", bg=C_HEADER_BG, bold=True, color="FFFFFF", h="center")
emp_atr = Counter(d[2] for d in dados if calc_status(d[6], d[8]) == "ATRASADO")
for i, (emp, cnt) in enumerate(emp_atr.most_common(), 12):
    dash_cell(wd, i, 2, emp, bg=C_ATRASADO_BG, color=C_ATRASADO_FG)
    dash_cell(wd, i, 3, cnt, bg=C_ATRASADO_BG, bold=True, color=C_ATRASADO_FG, h="center")

# Alertas (vence hoje)
alertas = [(d[2], d[4][:55]+"…" if len(d[4])>55 else d[4]) for d in dados if calc_status(d[6],d[8]) == "ALERTA"]
wd.merge_cells("E11:F11")
dash_cell(wd, 11, 5, "⚠ ALERTA — VENCE HOJE (20/05)", bg=C_ALERTA_BG, bold=True, color=C_ALERTA_FG, h="center")
for i, (emp, desc) in enumerate(alertas, 12):
    dash_cell(wd, i, 5, emp,  bg=C_ALERTA_BG, color=C_ALERTA_FG)
    c = wd.cell(row=i, column=6, value=desc)
    c.fill      = pf(C_ALERTA_BG)
    c.font      = ft(color=C_ALERTA_FG)
    c.border    = border_thin()
    c.alignment = al(h="left", wrap=True)
    wd.row_dimensions[i].height = 30

# ============================================================
# SALVAR
# ============================================================
output_path = "/home/user/meu-planejamento/Restricoes_Consolidadas_Cristalia_19052026_modificada_para_ReuniãodePlanejamentosS20.xlsx"
wb.save(output_path)
print(f"✅ Arquivo salvo: {output_path}")
print(f"   Total de registros: {len(dados)}")
print(f"   Status: {dict(Counter(status_list))}")
