#!/usr/bin/env python3
"""PUSH LIST REV1 — ASU Jundiaí (26001) | OTZ Engenharia × Messer Gases for Life"""

import os, math, base64, io, json
from datetime import date
from collections import defaultdict

try:
    from PIL import Image
    PIL_OK = True
except ImportError:
    PIL_OK = False

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT     = os.path.dirname(os.path.abspath(__file__))
UPLOAD   = "/root/.claude/uploads/001b7fb0-8857-54b0-9eb5-28053bfc3f61"
EXCEL_IN = os.path.join(UPLOAD, "da0cc769-DAENG_PUSH_LIST_ASU_JUNDIAI_REV0_CORRE__O.xlsx")
HTML_OUT = os.path.join(ROOT, "DAENG_PUSH_LIST_ASU_JUNDIAI.html")
XLSX_OUT = os.path.join(ROOT, "DAENG_PUSH_LIST_ASU_JUNDIAI.xlsx")
LOGO_OTZ = ("/tmp/claude-0/-home-user-meu-planejamento"
            "/001b7fb0-8857-54b0-9eb5-28053bfc3f61/scratchpad/logo_p1_img2_210x161.png")

LEVAN_DATE  = date(2026, 7, 2)
TODAY       = date.today()
DIAS_ABERTO = max(1, (TODAY - LEVAN_DATE).days)

# Renames confirmed by user
NOME_FIX = {
    "Escada 3 — Bacia de Contenção":  "Escada 3 — Portaria Definitiva do Projeto ASU",
    "Platôr Central":                  "Platô Central",
    "Platôr":                          "Platô Central",
}

DISC_COLOR = {
    "Macrodrenagem":      "#1565C0",
    "Terraplanagem":      "#6D4C41",
    "Plantio de Grama":   "#2E7D32",
    "Geral / 5S":         "#E65100",
    "Macrodrenagem / 5S": "#0277BD",
}
DEF_COLOR = "#546E7A"

# ── Area meta (GPS, horário, fotos originais) ─────────────────────────────────
AREA_META = {
    "01": {"gps":"23.1710S/46.9501W","horario":"15:14","data":"02/07/2026",
           "endereco":"Av. Pref. Luís Latorre, 9600 — Setor Industrial, Jundiaí/SP",
           "fotos":["0dce6895-1001308572.jpg"]},
    "02": {"gps":"23.1710S/46.9501W","horario":"15:16","data":"02/07/2026",
           "endereco":"Av. Pref. Luís Latorre, 9600 — Setor Industrial, Jundiaí/SP",
           "fotos":["f46133ce-1001308573.jpg"]},
    "03": {"gps":"23.1706S/46.9497W","horario":"15:34","data":"02/07/2026",
           "endereco":"Av. Pref. Luís Latorre, 9600 — Setor Industrial, Jundiaí/SP",
           "fotos":["ee4ec5be-1001308585.jpg"]},
    "04": {"gps":"23.1703S/46.9492W","horario":"15:51","data":"02/07/2026",
           "endereco":"Rod. Dom Gabriel Paulino Bueno Couto, 2045 — Chácara Terra Nova, Jundiaí/SP",
           "fotos":["cbac03ca-1001308613.jpg"]},
    "05": {"gps":"23.1703S/46.9492W","horario":"15:57","data":"02/07/2026",
           "endereco":"Av. Pref. Luís Latorre, 9600 — Setor Industrial, Jundiaí/SP",
           "fotos":["d8dad6bf-1001308617.jpg","c736ad32-1001308623.jpg","98a6888a-1001308624.jpg"]},
    "06": {"gps":"23.1703S/46.9492W","horario":"16:00","data":"02/07/2026",
           "endereco":"Av. Pref. Luís Latorre, 9600 — Setor Industrial, Jundiaí/SP",
           "fotos":["bd51de41-1001308618.jpg","37409c2b-1001308619.jpg"]},
    "07": {"gps":"23.1702S/46.9491W","horario":"16:06","data":"02/07/2026",
           "endereco":"Av. Pref. Luís Latorre, 9450–9600 — Setor Industrial, Jundiaí/SP",
           "fotos":["cb7d3a33-1001308635.jpg","840e9892-1001308642.jpg"]},
    "08": {"gps":"23.1702S/46.9483W","horario":"16:14","data":"02/07/2026",
           "endereco":"Rua Emílio Antonon, 1680 / Rua Nancy Carlota Netto, 120 — Jundiaí/SP",
           "fotos":["38482029-1001308650.jpg","f7549f10-1001308651.jpg",
                    "7d5f4cae-1001308652.jpg","8f08fbcf-1001308653.jpg"]},
    "09": {"gps":"23.1703S/46.9483W","horario":"16:17","data":"02/07/2026",
           "endereco":"Rua Emílio Antonon, 1680 / Av. Pref. Luís Latorre, 2067 — Jundiaí/SP",
           "fotos":["8997db98-1001308654.jpg","a393ade5-1001308655.jpg"]},
    "10": {"gps":"23.1703S/46.9483W","horario":"16:21","data":"02/07/2026",
           "endereco":"Rua Emílio Antonon, 1680 — Chácara Aeroporto, Jundiaí/SP",
           "fotos":["27b3328c-1001308656.jpg","64e95537-1001308657.jpg","b739367c-1001308658.jpg"]},
    "11": {"gps":"23.1702–23.1715S/46.9479–46.9483W","horario":"16:25–16:30","data":"02/07/2026",
           "endereco":"Av. Antonieta Piva Barranqueiros, 982 / Rua Nancy Carlota Netto, 120 — Jundiaí/SP",
           "fotos":["c8236fab-1001308670.jpg","eafa0b33-1001308671.jpg",
                    "3108b587-1001308672.jpg","0ce7bdaa-1001308679.jpg",
                    "bcb014dd-1001308680.jpg","f349934a-1001308681.jpg",
                    "ae810a7b-1001308682.jpg","1ce90304-1001308688.jpg",
                    "3bf04eef-1001308690.jpg","1a6375af-1001308691.jpg",
                    "9f14ec2e-1001308697.jpg","c504e614-1001308703.jpg"]},
    "12": {"gps":"23.1717S/46.9482W","horario":"16:37","data":"02/07/2026",
           "endereco":"Rod. Dom Gabriel Paulino Bueno Couto, 11010 — Medeiros, Jundiaí/SP",
           "fotos":["a3f56521-1001308704.jpg","106c1353-1001308705.jpg"]},
    "13": {"gps":"23.1718S/46.9481–46.9482W","horario":"16:42","data":"02/07/2026",
           "endereco":"Rua Emílio Antonon, 1680 — Chácara Aeroporto, Jundiaí/SP",
           "fotos":["76a9d794-1001308709.jpg","7f48bd14-1001308710.jpg","36220157-1001308711.jpg"]},
    "14": {"gps":"23.1721–23.1722S/46.9480W","horario":"16:45","data":"02/07/2026",
           "endereco":"Rua Emílio Antonon, 1680 / Rod. Dom Gabriel Paulino Bueno Couto, 11010 — Jundiaí/SP",
           "fotos":["f26293ac-1001308715.jpg","cfb019ad-1001308716.jpg",
                    "c7b1c31f-1001308717.jpg","c2fe493a-1001308720.jpg"]},
    "15": {"gps":"23.1721S/46.9488W","horario":"16:51","data":"02/07/2026",
           "endereco":"Rod. Dom Gabriel Paulino Bueno Couto, 11010 — Medeiros, Jundiaí/SP",
           "fotos":["7587ac67-1001308725.jpg","03e95a95-1001308726.jpg",
                    "c133effc-1001308727.jpg","a72461ba-1001308728.jpg"]},
    "16": {"gps":"23.1712S/46.9491W – 23.1710S/46.9494W","horario":"09:36","data":"03/07/2026",
           "endereco":"Rod. Dom Gabriel Paulino Bueno Couto, 11010 + Av. Pref. Luís Latorre, 9450 — Jundiaí/SP",
           "fotos":[]},
}

# ── Helpers ───────────────────────────────────────────────────────────────────

def fix_name(name):
    return NOME_FIX.get(name, name)

def disc_color(disc):
    return DISC_COLOR.get(disc, DEF_COLOR)

def encode_photo(fname, max_dim=700, quality=65):
    path = os.path.join(UPLOAD, fname)
    if not os.path.exists(path):
        return ""
    try:
        if PIL_OK:
            img = Image.open(path)
            img = img.convert("RGB")
            w, h = img.size
            if max(w, h) > max_dim:
                ratio = max_dim / max(w, h)
                img = img.resize((int(w*ratio), int(h*ratio)), Image.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=quality, optimize=True)
            return base64.b64encode(buf.getvalue()).decode()
        else:
            with open(path, "rb") as f:
                return base64.b64encode(f.read()).decode()
    except Exception as e:
        print(f"  [WARN] photo {fname}: {e}")
        return ""

def load_otz_logo():
    if not os.path.exists(LOGO_OTZ):
        return ""
    with open(LOGO_OTZ, "rb") as f:
        return base64.b64encode(f.read()).decode()

# ── Gauge SVG (240° arc, speedometer style) ──────────────────────────────────

def gauge_svg(pct, done_peso, total_peso, uid):
    W, H = 240, 175
    cx, cy, R = 120, 108, 78

    def polar(deg, r):
        rad = math.radians(deg)
        return (cx + r * math.cos(rad), cy - r * math.sin(rad))

    def arc_d(r, from_deg, to_deg):
        x0, y0 = polar(from_deg, r)
        x1, y1 = polar(to_deg, r)
        span = from_deg - to_deg
        large = 1 if span > 180 else 0
        return f"M{x0:.2f},{y0:.2f} A{r},{r} 0 {large},1 {x1:.2f},{y1:.2f}"

    SPAN, START = 240, 210
    end_angle = START - SPAN  # = -30
    prog_angle = START - (pct / 100 * SPAN)
    color = "#F44336" if pct < 33 else "#FFC107" if pct < 66 else "#00E676"

    lines = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
             f'width="{W}" height="{H}" style="display:block">']

    # Filters
    lines.append(f'''<defs>
  <filter id="glow{uid}" x="-40%" y="-40%" width="180%" height="180%">
    <feGaussianBlur in="SourceGraphic" stdDeviation="4" result="b"/>
    <feMerge><feMergeNode in="b"/><feMergeNode in="SourceGraphic"/></feMerge>
  </filter>
  <filter id="needleglow{uid}" x="-100%" y="-100%" width="300%" height="300%">
    <feGaussianBlur in="SourceGraphic" stdDeviation="2.5" result="b"/>
    <feMerge><feMergeNode in="b"/><feMergeNode in="SourceGraphic"/></feMerge>
  </filter>
  <linearGradient id="zoneg{uid}" gradientUnits="userSpaceOnUse"
    x1="{polar(210,R)[0]:.1f}" y1="{polar(210,R)[1]:.1f}"
    x2="{polar(-30,R)[0]:.1f}" y2="{polar(-30,R)[1]:.1f}">
    <stop offset="0%"   stop-color="#F44336"/>
    <stop offset="33%"  stop-color="#FF7043"/>
    <stop offset="66%"  stop-color="#FFC107"/>
    <stop offset="100%" stop-color="#00E676"/>
  </linearGradient>
</defs>''')

    # Background track (dark)
    lines.append(f'<path d="{arc_d(R, START, end_angle)}" fill="none" '
                 f'stroke="#1A3A60" stroke-width="18" stroke-linecap="round"/>')

    # Zone gradient track (subtle)
    lines.append(f'<path d="{arc_d(R, START, end_angle)}" fill="none" '
                 f'stroke="url(#zoneg{uid})" stroke-width="18" stroke-linecap="round" opacity="0.15"/>')

    # Progress arc
    if pct > 0.1:
        lines.append(f'<path d="{arc_d(R, START, prog_angle)}" fill="none" '
                     f'stroke="{color}" stroke-width="14" stroke-linecap="round" '
                     f'filter="url(#glow{uid})" opacity="0.95"/>')

    # Tick marks
    for tp in [0, 25, 50, 75, 100]:
        ta = START - (tp / 100 * SPAN)
        xi, yi = polar(ta, R - 10)
        xo, yo = polar(ta, R + 10)
        lw = "2.5" if tp in (0, 100) else "1.5"
        lines.append(f'<line x1="{xi:.1f}" y1="{yi:.1f}" x2="{xo:.1f}" y2="{yo:.1f}" '
                     f'stroke="#4A90D9" stroke-width="{lw}"/>')
        xl, yl = polar(ta, R + 24)
        lines.append(f'<text x="{xl:.1f}" y="{yl:.1f}" text-anchor="middle" '
                     f'dominant-baseline="middle" fill="#4A90D9" '
                     f'font-size="9" font-family="Arial,sans-serif">{tp}%</text>')

    # Minor ticks (every 10%)
    for tp in [10, 20, 30, 40, 60, 70, 80, 90]:
        ta = START - (tp / 100 * SPAN)
        xi, yi = polar(ta, R - 5)
        xo, yo = polar(ta, R + 5)
        lines.append(f'<line x1="{xi:.1f}" y1="{yi:.1f}" x2="{xo:.1f}" y2="{yo:.1f}" '
                     f'stroke="#2A4A70" stroke-width="1"/>')

    # Needle
    nx, ny = polar(prog_angle, R - 8)
    bx1, by1 = polar(prog_angle + 90, 6)
    bx2, by2 = polar(prog_angle - 90, 6)
    lines.append(f'<polygon points="{cx:.1f},{cy:.1f} {bx1:.1f},{by1:.1f} '
                 f'{nx:.1f},{ny:.1f} {bx2:.1f},{by2:.1f}" '
                 f'fill="white" filter="url(#needleglow{uid})" opacity="0.9"/>')

    # Hub rings
    lines.append(f'<circle cx="{cx}" cy="{cy}" r="9" fill="#0D1F3C" stroke="#00BCD4" stroke-width="2"/>')
    lines.append(f'<circle cx="{cx}" cy="{cy}" r="4" fill="#00BCD4"/>')

    # Percentage text
    if total_peso > 0:
        pct_str = f"{pct:.0f}%"
        lines.append(f'<text x="{cx}" y="{cy+30}" text-anchor="middle" '
                     f'fill="white" font-size="24" font-weight="bold" '
                     f'font-family="Arial,sans-serif" letter-spacing="1">{pct_str}</text>')
        lines.append(f'<text x="{cx}" y="{cy+46}" text-anchor="middle" '
                     f'fill="#00BCD4" font-size="8.5" font-family="Arial,sans-serif" '
                     f'letter-spacing="0.5">AVANÇO PONDERADO</text>')
        lines.append(f'<text x="{cx}" y="{cy+59}" text-anchor="middle" '
                     f'fill="#4A90D9" font-size="8" font-family="Arial,sans-serif">'
                     f'{done_peso}/{total_peso} pts</text>')
    else:
        lines.append(f'<text x="{cx}" y="{cy+35}" text-anchor="middle" '
                     f'fill="#4A90D9" font-size="11" font-family="Arial,sans-serif">S/ peso</text>')

    lines.append('</svg>')
    return "\n".join(lines)

# ── Read Excel ────────────────────────────────────────────────────────────────

def read_excel():
    wb = openpyxl.load_workbook(EXCEL_IN)
    ws = wb.active
    card_data = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        area      = fix_name(str(row[1]).strip()) if row[1] else ""
        local_num = str(row[2]).strip().zfill(2) if row[2] else "00"
        local_nom = fix_name(str(row[3]).strip()) if row[3] else ""
        foto_str  = str(row[4]).strip() if row[4] else "S/F"
        disc      = str(row[6]).strip() if row[6] else "Geral"
        desc      = str(row[7]).strip() if row[7] else ""
        peso      = int(row[8]) if row[8] and str(row[8]).strip().isdigit() else 1
        status    = str(row[9]).strip() if row[9] else "Pendente"
        prev_     = str(row[10]).strip() if row[10] else ""
        conc      = str(row[11]).strip() if row[11] else ""

        key = (local_num, foto_str)
        if key not in card_data:
            card_data[key] = {
                "area": area, "nome": local_nom,
                "local_num": local_num, "foto_str": foto_str,
                "atividades": []
            }
        card_data[key]["atividades"].append({
            "disciplina": disc, "descricao": desc,
            "peso": peso, "status": status,
            "previsao": prev_, "conclusao": conc
        })
    return card_data

# ── Build card list (ordered, with renumbered activities) ─────────────────────

def build_cards(card_data):
    # Group keys by local_num, sorted by foto_idx_num
    from_local = defaultdict(list)
    for (lnum, fstr) in card_data:
        if fstr == "S/F":
            fidx = 0
        else:
            try:
                fidx = int(fstr[1:3])
            except Exception:
                fidx = 0
        from_local[lnum].append((fidx, fstr))

    cards = []
    global_id = 0

    for lnum in sorted(from_local.keys()):
        entries = sorted(from_local[lnum], key=lambda x: x[0])
        total_remaining = sum(1 for fidx, _ in entries if fidx > 0)
        seq = 0

        for fidx, fstr in entries:
            key = (lnum, fstr)
            data = card_data[key]
            meta = AREA_META.get(lnum, {})

            # Photo file
            foto_b64 = ""
            sem_foto = True
            if fidx > 0:
                fotos = meta.get("fotos", [])
                if 0 < fidx <= len(fotos):
                    fname = fotos[fidx - 1]
                    print(f"  Encoding photo {fname} ...", end=" ", flush=True)
                    foto_b64 = encode_photo(fname)
                    print("ok" if foto_b64 else "FAIL")
                    sem_foto = False
                seq += 1
                foto_seq   = seq
                foto_total = total_remaining
            else:
                foto_seq   = 0
                foto_total = 0

            global_id += 1

            # Weighted progress
            atividades = []
            done_peso  = 0
            total_peso = 0
            for i, atv in enumerate(data["atividades"]):
                p = atv["peso"]
                total_peso += p
                if atv["status"].lower().startswith("conclu"):
                    done_peso += p
                atividades.append({**atv, "num": i + 1})

            pct = (done_peso / total_peso * 100) if total_peso > 0 else 0

            cards.append({
                "id":          global_id,
                "area":        data["area"],
                "nome":        data["nome"],
                "local_num":   lnum,
                "foto_seq":    foto_seq,
                "foto_total":  foto_total,
                "foto_b64":    foto_b64,
                "sem_foto":    sem_foto,
                "gps":         meta.get("gps", ""),
                "horario":     meta.get("horario", ""),
                "data_lev":    meta.get("data", ""),
                "endereco":    meta.get("endereco", ""),
                "atividades":  atividades,
                "done_peso":   done_peso,
                "total_peso":  total_peso,
                "pct":         pct,
            })
    return cards

# ── KPI computation ───────────────────────────────────────────────────────────

def compute_kpis(cards):
    tot_cards  = len(cards)
    tot_atv    = sum(len(c["atividades"]) for c in cards)
    tot_peso   = sum(c["total_peso"] for c in cards)
    done_atv   = sum(1 for c in cards for a in c["atividades"]
                     if a["status"].lower().startswith("conclu"))
    done_peso  = sum(c["done_peso"] for c in cards)
    pend_atv   = tot_atv - done_atv

    pct_simples   = (done_atv / tot_atv * 100) if tot_atv else 0
    pct_ponderado = (done_peso / tot_peso * 100) if tot_peso else 0

    # Área mais crítica (most pending activities)
    area_pending = defaultdict(int)
    for c in cards:
        for a in c["atividades"]:
            if not a["status"].lower().startswith("conclu"):
                area_pending[c["area"]] += 1
    area_critica = max(area_pending, key=area_pending.get) if area_pending else "—"
    area_critica_n = area_pending.get(area_critica, 0)

    # Disciplina em atraso (most pending)
    disc_pending = defaultdict(int)
    for c in cards:
        for a in c["atividades"]:
            if not a["status"].lower().startswith("conclu"):
                disc_pending[a["disciplina"]] += 1
    disc_atraso = max(disc_pending, key=disc_pending.get) if disc_pending else "—"
    disc_atraso_n = disc_pending.get(disc_atraso, 0)

    kpis = {
        "tot_cards": tot_cards,
        "tot_atv": tot_atv,
        "tot_peso": tot_peso,
        "done_atv": done_atv,
        "done_peso": done_peso,
        "pend_atv": pend_atv,
        "pct_simples": pct_simples,
        "pct_ponderado": pct_ponderado,
        "area_critica": area_critica,
        "area_critica_n": area_critica_n,
        "disc_atraso": disc_atraso,
        "disc_atraso_n": disc_atraso_n,
        "dias_aberto": DIAS_ABERTO,
        "velocidade": (done_atv / DIAS_ABERTO) if done_atv > 0 else None,
        "previsao": None,
    }

    if kpis["velocidade"] and pend_atv > 0:
        days_remaining = pend_atv / kpis["velocidade"]
        from datetime import timedelta
        eta = TODAY + timedelta(days=int(days_remaining))
        kpis["previsao"] = eta.strftime("%d/%m/%Y")

    return kpis

# ── HTML components ───────────────────────────────────────────────────────────

def html_kpi_box(icon, label, value, sub="", color="#00BCD4", always=True):
    if not always:
        return ""
    sub_html = f'<div class="kpi-sub">{sub}</div>' if sub else ""
    return f'''
<div class="kpi-box">
  <div class="kpi-icon" style="color:{color}">{icon}</div>
  <div class="kpi-val" style="color:{color}">{value}</div>
  <div class="kpi-lbl">{label}</div>
  {sub_html}
</div>'''

def html_activity_row(atv):
    color  = disc_color(atv["disciplina"])
    status = atv["status"]
    cls    = "concluido" if status.lower().startswith("conclu") else "pendente"
    prev   = atv["previsao"] or "—"
    conc   = atv["conclusao"] or "—"
    peso   = atv["peso"]
    disc   = atv["disciplina"]
    desc   = atv["descricao"]
    num    = atv["num"]

    return f'''<tr class="atv-row {cls}">
  <td class="atv-num">{num}</td>
  <td class="atv-disc" style="border-left:4px solid {color};padding-left:8px">
    <span class="disc-dot" style="background:{color}"></span>{disc}
  </td>
  <td class="atv-desc">{desc}</td>
  <td class="atv-peso"><span class="peso-badge" style="border-color:{color}">{peso}</span></td>
  <td class="atv-status {cls}">{status}</td>
  <td class="atv-date">{prev}</td>
  <td class="atv-date">{conc}</td>
</tr>'''

def html_card(card):
    uid    = str(card["id"])
    gauge  = gauge_svg(card["pct"], card["done_peso"], card["total_peso"], uid)

    if card["sem_foto"]:
        foto_html = '''<div class="photo-placeholder">
  <div class="photo-ph-icon">📷</div>
  <div class="photo-ph-txt">Foto não disponível</div>
</div>'''
        badge_html = ""
    else:
        foto_html = (f'<img class="card-photo" src="data:image/jpeg;base64,{card["foto_b64"]}" '
                     f'alt="Local {card["local_num"]} — foto {card["foto_seq"]}">')
        badge_html = (f'<div class="photo-badge">Foto {card["foto_seq"]} '
                      f'de {card["foto_total"]}</div>')

    atv_rows = "\n".join(html_activity_row(a) for a in card["atividades"])
    n_atv    = len(card["atividades"])
    done_cnt = sum(1 for a in card["atividades"] if a["status"].lower().startswith("conclu"))
    pend_cnt = n_atv - done_cnt
    pct_s    = f'{card["pct"]:.0f}'

    # Card area badge color based on local_num
    lnum_int = int(card["local_num"])
    hue = (lnum_int * 23) % 360
    badge_bg = f"hsl({hue},60%,30%)"
    badge_border = f"hsl({hue},60%,50%)"

    return f'''
<div class="push-card" id="card-{card["id"]}">
  <div class="card-header-bar" style="background:linear-gradient(135deg,{badge_bg},{badge_border}20)">
    <span class="card-id-badge">CARD {card["id"]:02d}</span>
    <span class="card-area-badge">Local {card["local_num"]}</span>
    <span class="card-nome">{card["nome"]}</span>
    <span class="card-pct-badge" style="color:{('#00E676' if card["pct"]>=66 else '#FFC107' if card["pct"]>=33 else '#F44336')}">{pct_s}%</span>
  </div>

  <div class="card-body">
    <div class="card-photo-col">
      <div class="photo-wrap">
        {foto_html}
        {badge_html}
      </div>
      <div class="card-meta-info">
        <div class="meta-row"><span class="meta-icon">◈</span><span class="meta-area">{card["area"]}</span></div>
        <div class="meta-row"><span class="meta-icon">⊕</span>{card["gps"]}</div>
        <div class="meta-row"><span class="meta-icon">◷</span>{card["horario"]} — {card["data_lev"]}</div>
        <div class="meta-row meta-addr"><span class="meta-icon">⊞</span>{card["endereco"]}</div>
      </div>
    </div>

    <div class="card-gauge-col">
      {gauge}
      <div class="card-stats">
        <span class="stat-badge done">{done_cnt} concl.</span>
        <span class="stat-badge pend">{pend_cnt} pend.</span>
        <span class="stat-badge total">{n_atv} total</span>
      </div>
    </div>
  </div>

  <div class="card-activities">
    <table class="atv-table">
      <thead>
        <tr>
          <th style="width:32px">#</th>
          <th>Disciplina</th>
          <th>Atividade / Descrição</th>
          <th style="width:44px">Peso</th>
          <th style="width:90px">Status</th>
          <th style="width:82px">Previsão</th>
          <th style="width:82px">Conclusão</th>
        </tr>
      </thead>
      <tbody>
        {atv_rows}
      </tbody>
    </table>
  </div>
</div>'''

# ── Main HTML assembly ────────────────────────────────────────────────────────

def generate_html(cards, kpis, otz_b64):
    today_str = TODAY.strftime("%d/%m/%Y")

    # OTZ logo
    if otz_b64:
        otz_img = f'<img src="data:image/png;base64,{otz_b64}" class="logo-otz" alt="OTZ / GPLAN">'
    else:
        otz_img = '<div class="logo-fallback">OTZ<br>GPLAN</div>'

    # Messer SVG
    messer_svg = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 220 68" class="logo-messer">
  <text x="0" y="44" font-family="'Arial Black',Arial,sans-serif" font-size="42"
        font-weight="900" fill="#1A3680">MESSER</text>
  <circle cx="196" cy="22" r="20" fill="none" stroke="#CC2020" stroke-width="2.5"/>
  <rect x="190" y="8" width="12" height="28" rx="3" fill="#CC2020"/>
  <rect x="183" y="15" width="26" height="14" rx="3" fill="#CC2020"/>
  <circle cx="196" cy="22" r="5.5" fill="white"/>
  <text x="0" y="64" font-family="Arial,sans-serif" font-size="15"
        fill="#CC2020" font-style="italic">Gases for Life</text>
</svg>'''

    # KPI section
    k = kpis
    kpi_html = "".join([
        html_kpi_box("⊞", "Locais / Cards",      str(k["tot_cards"]),         color="#00BCD4"),
        html_kpi_box("≡", "Total Atividades",    str(k["tot_atv"]),            color="#29B6F6"),
        html_kpi_box("⊙", "Total Peso",          f'{k["tot_peso"]} pts',       color="#4DD0E1"),
        html_kpi_box("✓", "Concluídas",          str(k["done_atv"]),           color="#00E676"),
        html_kpi_box("○", "Pendentes",           str(k["pend_atv"]),           color="#FFC107"),
        html_kpi_box("◎", "Avanço Simples",      f'{k["pct_simples"]:.1f}%',   color="#FF7043"),
        html_kpi_box("◉", "Avanço Ponderado",    f'{k["pct_ponderado"]:.1f}%', color="#E040FB"),
        html_kpi_box("▲", "Área Mais Crítica",   k["area_critica"],
                     sub=f'{k["area_critica_n"]} pendentes',                   color="#FF5252"),
        html_kpi_box("◆", "Disciplina em Atraso", k["disc_atraso"],
                     sub=f'{k["disc_atraso_n"]} pendentes',                   color="#FF6D00"),
        html_kpi_box("⏱", "Dias em Aberto",      str(k["dias_aberto"]),
                     sub=f'desde {LEVAN_DATE.strftime("%d/%m/%Y")}',          color="#B0BEC5"),
        html_kpi_box("⚡", "Velocidade",
                     f'{k["velocidade"]:.2f} atv/dia' if k["velocidade"] else "—",
                     sub="concluídas/dia",
                     color="#69F0AE", always=(k["velocidade"] is not None)),
        html_kpi_box("⊗", "Previsão de Conclusão", k["previsao"] or "—",
                     color="#CE93D8", always=(k["previsao"] is not None)),
    ])

    cards_html = "\n".join(html_card(c) for c in cards)

    return f'''<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>PUSH LIST — DAENG | ASU Jundiaí (26001) — REV1</title>
<style>
/* ── Reset & base ─────────────────────────────────────────────── */
*,*::before,*::after{{box-sizing:border-box;margin:0;padding:0}}
body{{
  background:#0A1628;color:#C8D8F0;
  font-family:'Segoe UI',Arial,sans-serif;font-size:13px;line-height:1.4;
  min-height:100vh;
}}
h1,h2,h3{{font-weight:700;letter-spacing:.5px}}

/* ── Header ───────────────────────────────────────────────────── */
.site-header{{
  background:linear-gradient(135deg,#0A1628 0%,#0D2540 40%,#0A1628 100%);
  border-bottom:2px solid #00BCD4;
  padding:0;
}}
.header-top{{
  display:flex;align-items:center;gap:16px;
  padding:14px 20px;
}}
.logo-otz{{height:64px;width:auto;filter:drop-shadow(0 0 6px #00BCD440)}}
.logo-fallback{{
  background:#0D2540;border:1px solid #00BCD4;border-radius:6px;
  padding:6px 12px;font-size:11px;font-weight:700;color:#00BCD4;text-align:center;
}}
.logo-messer{{height:52px;width:auto}}
.header-title{{flex:1;text-align:center}}
.header-title h1{{
  font-size:20px;color:#00BCD4;letter-spacing:1.5px;text-transform:uppercase;
  text-shadow:0 0 16px #00BCD480;
}}
.header-title .sub{{font-size:12px;color:#7BAFD4;margin-top:3px}}
.header-refbar{{
  background:#06101E;border-top:1px solid #1A3A60;
  padding:8px 20px;font-size:10.5px;color:#4A90D9;
  display:flex;flex-wrap:wrap;gap:8px;align-items:center;
}}
.refbar-pill{{
  background:#0D1F3C;border:1px solid #1A3A60;border-radius:4px;
  padding:2px 8px;white-space:nowrap;
}}
.refbar-pill.doc{{border-color:#00BCD440;color:#29B6F6;font-weight:600}}

/* ── KPI section ──────────────────────────────────────────────── */
.kpi-section{{
  background:linear-gradient(180deg,#06101E,#0A1628);
  border-bottom:1px solid #1A3A60;padding:16px 20px;
}}
.kpi-title{{font-size:11px;color:#4A90D9;letter-spacing:2px;
           text-transform:uppercase;margin-bottom:12px}}
.kpi-grid{{
  display:flex;flex-wrap:wrap;gap:10px;
}}
.kpi-box{{
  background:#0D1F3C;border:1px solid #1A3A60;border-radius:8px;
  padding:10px 14px;min-width:120px;flex:1;text-align:center;
  transition:transform .2s,box-shadow .2s;
}}
.kpi-box:hover{{transform:translateY(-2px);box-shadow:0 4px 16px #00BCD420}}
.kpi-icon{{font-size:18px;margin-bottom:4px}}
.kpi-val{{font-size:20px;font-weight:700;line-height:1.1}}
.kpi-lbl{{font-size:10px;color:#7BAFD4;margin-top:3px;text-transform:uppercase;letter-spacing:.8px}}
.kpi-sub{{font-size:10px;color:#4A90D9;margin-top:2px}}

/* ── Cards grid ───────────────────────────────────────────────── */
.cards-section{{padding:16px 20px}}
.cards-section-title{{
  font-size:11px;color:#4A90D9;letter-spacing:2px;text-transform:uppercase;
  margin-bottom:14px;display:flex;align-items:center;gap:8px;
}}
.cards-section-title::after{{
  content:'';flex:1;height:1px;background:linear-gradient(90deg,#1A3A60,transparent);
}}
.cards-grid{{
  display:grid;
  grid-template-columns:repeat(auto-fill,minmax(700px,1fr));
  gap:16px;
}}

/* ── Individual push card ─────────────────────────────────────── */
.push-card{{
  background:#0D1F3C;border:1px solid #1A3A60;border-radius:10px;
  overflow:hidden;transition:box-shadow .25s;
}}
.push-card:hover{{box-shadow:0 4px 24px #00BCD430}}
.card-header-bar{{
  display:flex;align-items:center;gap:10px;padding:8px 14px;
  border-bottom:1px solid #1A3A60;
}}
.card-id-badge{{
  background:#00BCD4;color:#0A1628;font-size:11px;font-weight:700;
  border-radius:4px;padding:2px 8px;white-space:nowrap;letter-spacing:.5px;
}}
.card-area-badge{{
  background:#1A3A60;color:#7BAFD4;font-size:11px;
  border-radius:4px;padding:2px 8px;white-space:nowrap;
}}
.card-nome{{flex:1;font-size:13px;font-weight:600;color:#C8D8F0}}
.card-pct-badge{{font-size:16px;font-weight:700;white-space:nowrap}}

/* ── Card body (photo + gauge) ────────────────────────────────── */
.card-body{{
  display:flex;gap:0;border-bottom:1px solid #1A3A60;
}}
.card-photo-col{{
  width:260px;min-width:220px;border-right:1px solid #1A3A60;
  display:flex;flex-direction:column;
}}
.photo-wrap{{position:relative;height:195px;background:#050F1A;overflow:hidden}}
.card-photo{{width:100%;height:100%;object-fit:cover;display:block}}
.photo-placeholder{{
  width:100%;height:100%;display:flex;flex-direction:column;
  align-items:center;justify-content:center;gap:8px;
  background:#060E1C;color:#2A4A70;
}}
.photo-ph-icon{{font-size:32px;opacity:.5}}
.photo-ph-txt{{font-size:11px}}
.photo-badge{{
  position:absolute;bottom:6px;right:6px;
  background:#00000080;color:#00BCD4;font-size:10px;
  border-radius:4px;padding:2px 6px;backdrop-filter:blur(4px);
}}
.card-meta-info{{
  padding:8px 10px;font-size:10.5px;color:#7BAFD4;
  display:flex;flex-direction:column;gap:4px;flex:1;
}}
.meta-row{{display:flex;align-items:flex-start;gap:5px;}}
.meta-icon{{color:#00BCD4;min-width:14px;}}
.meta-area{{color:#C8D8F0;font-weight:600;}}
.meta-addr{{font-size:10px;color:#4A90D9;line-height:1.3;}}

.card-gauge-col{{
  flex:1;display:flex;flex-direction:column;align-items:center;
  justify-content:center;padding:8px;
}}
.card-stats{{
  display:flex;gap:6px;margin-top:4px;
}}
.stat-badge{{
  font-size:10px;border-radius:4px;padding:2px 7px;font-weight:600;
}}
.stat-badge.done{{background:#003020;color:#00E676;border:1px solid #00E67640}}
.stat-badge.pend{{background:#302000;color:#FFC107;border:1px solid #FFC10740}}
.stat-badge.total{{background:#001830;color:#29B6F6;border:1px solid #29B6F640}}

/* ── Activity table ───────────────────────────────────────────── */
.card-activities{{overflow-x:auto}}
.atv-table{{
  width:100%;border-collapse:collapse;font-size:12px;
}}
.atv-table thead tr{{
  background:#06101E;border-bottom:2px solid #00BCD440;
}}
.atv-table th{{
  padding:6px 8px;color:#4A90D9;font-size:10.5px;
  font-weight:600;text-transform:uppercase;letter-spacing:.5px;text-align:left;
}}
.atv-table tbody tr{{border-bottom:1px solid #0F2040}}
.atv-table tbody tr:hover{{background:#0A2040}}
.atv-row.concluido{{background:#020A04}}
.atv-row.concluido:hover{{background:#061208}}
.atv-num{{
  padding:7px 8px;color:#4A90D9;font-weight:600;
  text-align:center;white-space:nowrap;
}}
.atv-disc{{
  padding:7px 8px;color:#C8D8F0;font-size:11.5px;white-space:nowrap;
}}
.disc-dot{{
  display:inline-block;width:8px;height:8px;border-radius:50%;
  margin-right:5px;vertical-align:middle;
}}
.atv-desc{{padding:7px 8px;color:#B0C8E8;line-height:1.35}}
.atv-peso{{padding:7px 6px;text-align:center}}
.peso-badge{{
  display:inline-block;width:22px;height:22px;line-height:22px;
  text-align:center;border-radius:50%;font-weight:700;font-size:11px;
  border:2px solid;color:#C8D8F0;
}}
.atv-status{{padding:7px 8px;font-size:11px;font-weight:600;white-space:nowrap}}
.atv-status.pendente{{color:#FFC107}}
.atv-status.concluido{{color:#00E676}}
.atv-date{{padding:7px 8px;font-size:11px;color:#4A90D9;white-space:nowrap}}

/* ── Footer ───────────────────────────────────────────────────── */
.site-footer{{
  background:#06101E;border-top:2px solid #00BCD430;
  padding:14px 20px;display:flex;justify-content:space-between;
  align-items:center;flex-wrap:wrap;gap:8px;font-size:11px;color:#4A90D9;
}}
.footer-by{{color:#7BAFD4;font-style:italic}}
.footer-rev{{
  background:#0D1F3C;border:1px solid #1A3A60;border-radius:4px;
  padding:3px 10px;font-weight:600;color:#00BCD4;
}}

/* ── Scrollbar ────────────────────────────────────────────────── */
::-webkit-scrollbar{{width:6px;height:6px}}
::-webkit-scrollbar-track{{background:#060E1C}}
::-webkit-scrollbar-thumb{{background:#1A3A60;border-radius:3px}}
::-webkit-scrollbar-thumb:hover{{background:#00BCD4}}

@media(max-width:760px){{
  .cards-grid{{grid-template-columns:1fr}}
  .card-body{{flex-direction:column}}
  .card-photo-col{{width:100%;border-right:none;border-bottom:1px solid #1A3A60}}
  .photo-wrap{{height:180px}}
  .kpi-grid{{gap:8px}}
  .kpi-box{{min-width:90px}}
}}
</style>
</head>
<body>

<!-- ── Header ─────────────────────────────────────────────────────────── -->
<header class="site-header">
  <div class="header-top">
    {otz_img}
    <div class="header-title">
      <h1>Push List — Serviços DAENG</h1>
      <div class="sub">ASU Jundiaí (26001) &nbsp;|&nbsp; OTZ Engenharia × Messer Gases for Life</div>
    </div>
    {messer_svg}
  </div>
  <div class="header-refbar">
    <span class="refbar-pill doc">P-DE-E-179-CIV18-000-003 — AIR SEPARATION UNIT MASTER PLAN | Terraplanagem — Planta | OTZ PGE | Aprov.: RNW | 19/08/25</span>
    <span class="refbar-pill">Ref: CLM-216</span>
    <span class="refbar-pill">Contratada: Andrade e Rocha (DAENG)</span>
    <span class="refbar-pill">Gerenciadora: OTZ Engenharia — GPLAN</span>
    <span class="refbar-pill">Levantamento: {LEVAN_DATE.strftime("%d/%m/%Y")}</span>
    <span class="refbar-pill">Elaborado por: Caio Sergio Bento da Silva &amp; Ilson dos Santos Azevedo</span>
    <span class="refbar-pill" style="color:#00BCD4;font-weight:600">REV1 — {today_str}</span>
  </div>
</header>

<!-- ── KPIs ───────────────────────────────────────────────────────────── -->
<section class="kpi-section">
  <div class="kpi-title">&#9670; Indicadores Gerais</div>
  <div class="kpi-grid">
    {kpi_html}
  </div>
</section>

<!-- ── Cards ──────────────────────────────────────────────────────────── -->
<section class="cards-section">
  <div class="cards-section-title">&#9670; Locais de Serviço &nbsp; ({len(cards)} cards &nbsp;·&nbsp; {kpis["tot_atv"]} atividades)</div>
  <div class="cards-grid">
    {cards_html}
  </div>
</section>

<!-- ── Footer ─────────────────────────────────────────────────────────── -->
<footer class="site-footer">
  <div class="footer-by">Elaborado por: Ilson do Santos Azevedo / Supervisor de Planejamento</div>
  <div>Projeto ASU Jundiaí (26001) &nbsp;|&nbsp; OTZ Engenharia × Messer Gases for Life &nbsp;|&nbsp; DAENG — Andrade e Rocha</div>
  <div class="footer-rev">REV1 &nbsp;·&nbsp; {today_str}</div>
</footer>

</body>
</html>'''

# ── Excel generation ──────────────────────────────────────────────────────────

def generate_xlsx(cards):
    wb   = openpyxl.Workbook()
    ws   = wb.active
    ws.title = "Push List REV1"

    # Colors
    C_HEADER  = "0A1628"
    C_SUBHDR  = "0D2540"
    C_ROW_ODD = "0D1A30"
    C_ROW_EVN = "081020"
    C_ACCENT  = "00BCD4"
    C_TEXT    = "C8D8F0"

    def style_cell(cell, bold=False, bg=None, fg="C8D8F0", sz=11, wrap=False,
                   h_align="left", v_align="center"):
        cell.font      = Font(bold=bold, color=fg, size=sz, name="Calibri")
        cell.alignment = Alignment(horizontal=h_align, vertical=v_align, wrap_text=wrap)
        if bg:
            cell.fill  = PatternFill("solid", fgColor=bg)

    thin  = Side(style="thin",   color="1A3A60")
    med   = Side(style="medium", color="00BCD4")
    brd   = Border(left=thin, right=thin, top=thin, bottom=thin)
    brd_m = Border(left=med,  right=med,  top=med,  bottom=med)

    # Row 1: Title
    ws.merge_cells("A1:M1")
    c = ws["A1"]
    c.value = "PUSH LIST — SERVIÇOS DAENG | ASU JUNDIAÍ (26001) | OTZ Engenharia × Messer Gases for Life"
    style_cell(c, bold=True, bg=C_HEADER, fg=C_ACCENT, sz=13, h_align="center")

    # Row 2: Subtitle
    ws.merge_cells("A2:M2")
    c = ws["A2"]
    c.value = (f"Ref: CLM-216  |  Contratada: Andrade e Rocha (DAENG)  |  "
               f"Gerenciadora: OTZ Engenharia — GPLAN  |  "
               f"Levantamento: {LEVAN_DATE.strftime('%d/%m/%Y')}  |  "
               f"REV1 — {TODAY.strftime('%d/%m/%Y')}  |  "
               f"Elaborado por: Ilson do Santos Azevedo / Supervisor de Planejamento")
    style_cell(c, bg=C_SUBHDR, fg="7BAFD4", sz=9, h_align="center")

    # Row 3: Headers
    headers = ["#Card", "Área", "Local Nº", "Nome do Local",
               "Foto", "# Ativ.", "Disciplina", "Atividade / Descrição",
               "Peso", "Status", "Previsão", "Conclusão", "Avanço (%)"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=h)
        style_cell(c, bold=True, bg=C_SUBHDR, fg=C_ACCENT, sz=10, h_align="center")
        c.border = brd_m

    # Column widths
    widths = [8, 30, 10, 34, 8, 8, 22, 52, 7, 12, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze panes
    ws.freeze_panes = "A4"

    row_num = 4
    for card in cards:
        n_atv  = len(card["atividades"])
        r_start = row_num
        for i, atv in enumerate(card["atividades"]):
            bg = C_ROW_ODD if (row_num % 2) == 0 else C_ROW_EVN
            vals = [
                card["id"], card["area"], card["local_num"], card["nome"],
                card["foto_seq"] if card["foto_seq"] > 0 else "S/F",
                atv["num"], atv["disciplina"], atv["descricao"],
                atv["peso"], atv["status"], atv["previsao"] or "",
                atv["conclusao"] or "", f'{card["pct"]:.1f}%'
            ]
            for col, val in enumerate(vals, 1):
                c = ws.cell(row=row_num, column=col, value=val)
                sz  = 10
                wrp = (col == 8)
                hal = "center" if col in (1, 3, 5, 6, 9, 10, 11, 12, 13) else "left"
                style_cell(c, bg=bg, sz=sz, wrap=wrp, h_align=hal)
                c.border = brd
            ws.row_dimensions[row_num].height = 16 if not atv["descricao"] or len(atv["descricao"]) < 70 else 30
            row_num += 1

        # Merge card-level cells across activity rows
        if n_atv > 1:
            for col in (1, 2, 3, 4, 5, 13):
                try:
                    ws.merge_cells(start_row=r_start, start_column=col,
                                   end_row=row_num-1, end_column=col)
                    ws.cell(r_start, col).alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=(col==2 or col==4))
                except Exception:
                    pass

    wb.save(XLSX_OUT)
    print(f"  Excel saved: {XLSX_OUT} ({row_num - 4} activity rows)")

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    print("=" * 60)
    print("PUSH LIST REV1 — ASU Jundiaí (26001)")
    print("=" * 60)

    print("\n[1] Reading Excel …")
    card_data = read_excel()
    print(f"    {len(card_data)} card keys loaded")

    print("\n[2] Building cards + encoding photos …")
    cards = build_cards(card_data)
    print(f"    {len(cards)} cards built")

    print("\n[3] Computing KPIs …")
    kpis = compute_kpis(cards)
    print(f"    Avanço ponderado: {kpis['pct_ponderado']:.1f}%")
    print(f"    Área crítica: {kpis['area_critica']}")
    print(f"    Disciplina em atraso: {kpis['disc_atraso']}")

    print("\n[4] Loading OTZ logo …")
    otz_b64 = load_otz_logo()
    print(f"    Logo: {'OK' if otz_b64 else 'not found (text fallback)'}")

    print("\n[5] Generating HTML …")
    html = generate_html(cards, kpis, otz_b64)
    with open(HTML_OUT, "w", encoding="utf-8") as f:
        f.write(html)
    size_mb = os.path.getsize(HTML_OUT) / (1024 * 1024)
    print(f"    HTML saved: {HTML_OUT} ({size_mb:.1f} MB)")

    print("\n[6] Generating Excel …")
    generate_xlsx(cards)

    print("\n✓ Done — REV1 complete")
    print(f"  HTML: {HTML_OUT}")
    print(f"  XLSX: {XLSX_OUT}")

if __name__ == "__main__":
    main()
